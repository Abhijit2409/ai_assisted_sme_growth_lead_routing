"""
Phase C — ISED / Statistics Canada Survey on Financing and Growth of SMEs, 2023
(data-tables workbook) — structural inventory and source validation.

Question: what does the workbook measure, which tables and dimensions does it
contain, how must its estimates be read, and can any table support the SME
market-opportunity analysis?

Design (mirrors 03/04):
- The workbook in data/raw is read-only: never resaved. Its sha256 is checked
  against data/source_manifest.csv before and after processing.
- The workbook is read twice with openpyxl (formulas preserved, cached values) and
  the two reads are compared cell by cell. Only cells that exist in the sheet XML
  are read (ws._cells); ws.cell() is never called, because it creates cells and
  widens the reported used range.
- Every table sheet is parsed generically: title/questions/notes, multi-row
  headers (merged and unmerged), the unit row (value / Quality / SE-or-CV
  triplets), the row stub and footnotes. A tidy cell table is written to a
  gitignored cache for later phases; compact inventories go to
  data/validation/ised_2023/.
- Interpretation that needs judgement (denominators, reconciliation specs,
  relevance) is declared explicitly in this file, with its evidence.
- Every numeric claim in the Phase C section of docs/02_data_sources_and_validation.md
  is checked against the outputs. Any failed gate exits with status 1.

Run (Anaconda interpreter):
    C:/Users/abhij/anaconda3/python.exe src/validation/00_build_source_manifest.py
    C:/Users/abhij/anaconda3/python.exe src/validation/05_validate_ised_sfgsme_2023.py
"""

from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime
import hashlib
import json
import re
import sys
import time
import unicodedata
import zipfile

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]   # project root, independent of the working directory

RAW_FILE = ROOT / "data" / "raw" / "ised_sme_financing_2023" / "sfgsme-efcpme-tables-2023-eng.xlsx"
MANIFEST_FILE = ROOT / "data" / "source_manifest.csv"
DOC_FILE = ROOT / "docs" / "02_data_sources_and_validation.md"
DOC_SECTION_HEADING = "## ISED Survey on Financing and Growth of SMEs 2023 (Phase C)"

CACHE_DIR = ROOT / "data" / "cache"
TIDY_CACHE = CACHE_DIR / "ised_sfgsme_2023_cells.parquet"
OUTPUT_DIR = ROOT / "data" / "validation" / "ised_2023"

# Phase A / B outputs used only for label-level compatibility diagnostics.
GEO_0722 = ROOT / "data" / "validation" / "33100722" / "geography_classification.csv"
GEO_1097 = ROOT / "data" / "validation" / "33101097" / "geography_profile.csv"

# Provenance verified on official pages (see docs). Portal file size is used as a gate.
PORTAL_RESOURCE_URL = ("https://open.canada.ca/data/en/dataset/c3c9f361-c7d3-4c87-8be0-64c2b62521ee/"
                       "resource/56f10b18-2aeb-4734-b0ad-8450bdf25d08")
PORTAL_FILE_BYTES = 2_482_697
IMDB_SAMPLE_SIZE = 29_100          # "final unduplicated size of the sample was 29,100 enterprises"

# Quality-code thresholds as published by the two official sources.
# Workbook "Notes" sheet: E = CV 20% up to 30% / SE 10 up to 15; F = CV 30%+ / SE 15+.
# IMDB record 2941:       E = CV 20% up to 25% / SE 10 up to 12.5; F = CV 25%+ / SE 12.5+.
THRESHOLDS = {
    "workbook_notes": {"CV": [5, 10, 15, 20, 30], "SE": [2.5, 5, 7.5, 10, 15]},
    "imdb_record":    {"CV": [5, 10, 15, 20, 25], "SE": [2.5, 5, 7.5, 10, 12.5]},
}
CODES = ["A", "B", "C", "D", "E", "F"]

VALUE_TOKENS = {"%", "$", "#", "Months", "Average", "average", "Median"}
QUALITY_TOKEN = "Quality"
ERROR_TOKENS = {"SE", "CV"}
HEADER_ROWS = [11, 12, 13, 14, 15]
FIRST_DATA_COL = 4                     # column D; A-C hold the row stub

gates = {}


def banner(title):
    print("\n" + "=" * 90)
    print(title)
    print("=" * 90)


def gate(name, passed, detail=""):
    gates[name] = bool(passed)
    print(f"  [{'PASS' if passed else 'FAIL'}] {name}" + (f"  ({detail})" if detail != "" else ""))


def sha256_file(path, chunk=1 << 20):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while block := f.read(chunk):
            h.update(block)
    return h.hexdigest()


def clean(v):
    """Collapse whitespace in a label; the raw workbook labels carry trailing newlines."""
    return re.sub(r"\s+", " ", str(v)).strip()


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def col_letter(c):
    return openpyxl.utils.get_column_letter(c)


def write_csv(df, name):
    """Write to OUTPUT_DIR, keeping integral float columns (NaN-padded counts) as integers."""
    df = df.copy()
    for c in df.columns:
        if df[c].dtype.kind == "f":
            s = df[c].dropna()
            if len(s) and (s == s.round()).all():
                df[c] = df[c].astype("Int64")
    df.to_csv(OUTPUT_DIR / name, index=False)


# ---------------------------------------------------
# 1. Package-level inspection (the xlsx is a zip of XML parts)
# ---------------------------------------------------

def inspect_package():
    z = zipfile.ZipFile(RAW_FILE)
    names = z.namelist()
    wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid_target = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    rid_target.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})

    sheets = []
    for m in re.finditer(r"<sheet ([^>]*)/>", wb_xml):
        attrs = dict(re.findall(r'([\w:]+)="([^"]*)"', m.group(1)))
        part = "xl/" + rid_target[attrs["r:id"]].lstrip("/").removeprefix("xl/")
        head = z.read(part)[:4000].decode("utf-8", "ignore")
        dim = re.search(r'<dimension ref="([^"]+)"', head)
        sheets.append({"name": attrs["name"], "state_xml": attrs.get("state", "visible"),
                       "xml_part": part, "xml_dimension": dim.group(1) if dim else None})

    defined = re.findall(r"<definedName ([^>]*)>(.*?)</definedName>", wb_xml)
    def_rows = []
    for attrs, ref in defined:
        a = dict(re.findall(r'([\w:]+)="([^"]*)"', attrs))
        name = a["name"]
        kind = ("autofilter_range" if name == "_xlnm._FilterDatabase"
                else "custom_view_filter_data" if re.match(r"Z_[0-9A-F_]+_\.wvu\.FilterData", name)
                else "print_area" if name == "_xlnm.Print_Area"
                else "print_titles" if name == "_xlnm.Print_Titles"
                else "custom_view_other" if ".wvu." in name
                else "user_named_range")
        def_rows.append({"name": name, "local_sheet_id": a.get("localSheetId"), "hidden": a.get("hidden") == "1",
                         "kind": kind, "refers_to": ref})

    app = z.read("docProps/app.xml").decode("utf-8")
    core = z.read("docProps/core.xml").decode("utf-8")
    titles_of_parts = re.search(r"<TitlesOfParts><vt:vector size=\"(\d+)\"", app)
    return {
        "zip_parts": len(names),
        "sheets": sheets,
        "defined_names": def_rows,
        "custom_workbook_views": wb_xml.count("<customWorkbookView "),
        "external_link_parts": sum(n.startswith("xl/externalLinks/") for n in names),
        "pivot_parts": sum("pivot" in n.lower() for n in names),
        "chart_parts": sum(n.startswith("xl/charts/") for n in names),
        "table_parts": sum(n.startswith("xl/tables/") for n in names),
        "comment_parts": sum("comments" in n.lower() for n in names),
        "drawing_parts": sum(n.startswith("xl/drawings/") and n.endswith(".xml") for n in names),
        "media_parts": sum(n.startswith("xl/media/") for n in names),
        "vba_parts": sum(n.lower().endswith(".bin") and "vba" in n.lower() for n in names),
        "app_titles_of_parts": int(titles_of_parts.group(1)) if titles_of_parts else None,
        "app_company": (re.search(r"<Company>(.*?)</Company>", app) or [None, None])[1],
        "core_creator": (re.search(r"<dc:creator>(.*?)</dc:creator>", core) or [None, None])[1],
        "core_last_modified_by": (re.search(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core) or [None, None])[1],
        "core_created": (re.search(r"<dcterms:created[^>]*>(.*?)</dcterms:created>", core) or [None, None])[1],
        "core_modified": (re.search(r"<dcterms:modified[^>]*>(.*?)</dcterms:modified>", core) or [None, None])[1],
    }


# ---------------------------------------------------
# 2. Sheet access helpers (existing cells only)
# ---------------------------------------------------

class Sheet:
    """Read-only view of existing cells; never creates cells."""

    def __init__(self, ws):
        self.ws = ws
        self.title = ws.title
        self.used_range = ws.calculate_dimension()      # computed before any access
        self.cells = {k: c for k, c in ws._cells.items()}
        self.nonempty = {k: c for k, c in self.cells.items() if c.value is not None}
        self.merged = list(ws.merged_cells.ranges)
        self.merge_map = {}
        for mr in self.merged:
            tl = self.cells.get((mr.min_row, mr.min_col))
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    self.merge_map[(r, c)] = (mr, tl.value if tl is not None else None)

    def value(self, r, c):
        cell = self.cells.get((r, c))
        return cell.value if cell is not None else None

    def header_value(self, r, c):
        """(value, is_segment_start) honouring merged ranges."""
        if (r, c) in self.merge_map:
            mr, v = self.merge_map[(r, c)]
            return v, c == mr.min_col
        v = self.value(r, c)
        return v, v is not None

    def bbox(self):
        if not self.nonempty:
            return None
        rows = [r for r, _ in self.nonempty]
        cols = [c for _, c in self.nonempty]
        return (min(rows), min(cols), max(rows), max(cols))


def render_label(cell):
    """Stub labels: numbers formatted as % (e.g. female-ownership 0 / 0.5 / 1) are rendered as displayed."""
    v = cell.value
    if is_number(v) and "%" in (cell.number_format or ""):
        return f"{round(v * 100, 6):g}%", True
    return clean(v), False


# ---------------------------------------------------
# 3. Table parsing
# ---------------------------------------------------

def parse_table(sh):
    m = re.match(r"Table\s+(\w+)\s*:\s*(.*)", clean(sh.value(1, 1) or ""))
    table_no = m.group(1) if m else None
    title = m.group(2) if m else None

    # Unit row: >= 2 unit tokens in the data columns.
    by_row = defaultdict(list)
    for (r, c), cell in sh.nonempty.items():
        if c >= FIRST_DATA_COL and isinstance(cell.value, str):
            by_row[r].append(cell.value.strip())
    unit_rows = sorted(r for r, toks in by_row.items()
                       if sum(t in VALUE_TOKENS | ERROR_TOKENS | {QUALITY_TOKEN, "Median"} for t in toks) >= 2)
    unit_row = unit_rows[0] if unit_rows else None

    source_rows = sorted(r for (r, c), cell in sh.nonempty.items()
                         if c == 1 and isinstance(cell.value, str) and cell.value.strip().startswith("Source"))
    source_row = source_rows[0] if source_rows else None
    last_col = max(c for (_, c) in sh.nonempty)

    # Column header paths (rows 11-15), honouring merged ranges and resetting lower rows at new segments.
    current = {r: None for r in HEADER_ROWS}
    paths, top_starts = {}, {}
    top_row = min((r for r in HEADER_ROWS for c in range(FIRST_DATA_COL, last_col + 1)
                   if sh.header_value(r, c)[0] not in (None, "")), default=None)
    for c in range(FIRST_DATA_COL, last_col + 1):
        for r in HEADER_ROWS:
            v, start = sh.header_value(r, c)
            if v not in (None, "") and start:
                current[r] = clean(v)
                for rr in HEADER_ROWS:
                    if rr > r:
                        current[rr] = None
                if r == top_row:
                    top_starts[c] = True
            elif v not in (None, ""):
                current[r] = clean(v)
        path = [current[r] for r in HEADER_ROWS if current[r]]
        # A header merged down several rows repeats; keep one copy of consecutive duplicates.
        paths[c] = [v for i, v in enumerate(path) if i == 0 or v != path[i - 1]]

    # Measures from the unit row: value column + attached Quality + SE/CV.
    measures, cur = [], None
    block_no, block_title = 0, None
    for c in range(FIRST_DATA_COL, last_col + 1):
        if top_starts.get(c):
            block_no += 1
            block_title = paths[c][0] if paths[c] else None
        tok = sh.value(unit_row, c) if unit_row else None
        tok = tok.strip() if isinstance(tok, str) else tok
        if tok in VALUE_TOKENS:
            cur = {"value_col": c, "unit_token": tok, "quality_col": None, "error_col": None, "error_type": None,
                   "header_path": paths[c], "block_no": max(block_no, 1), "block_title": block_title}
            measures.append(cur)
        elif tok == QUALITY_TOKEN and cur is not None and cur["quality_col"] is None:
            cur["quality_col"] = c
        elif tok in ERROR_TOKENS and cur is not None and cur["error_col"] is None:
            cur["error_col"], cur["error_type"] = c, tok

    # Row stub between the unit row and the Source footnote.
    stub, group = [], None
    for r in range(unit_row + 1, source_row):
        a, b, c = sh.cells.get((r, 1)), sh.cells.get((r, 2)), sh.cells.get((r, 3))
        if a is not None and a.value is not None:
            group = clean(a.value)
        lab_b = render_label(b) if b is not None and b.value is not None else (None, False)
        lab_c = render_label(c) if c is not None and c.value is not None else (None, False)
        has_values = any((r, cc) in sh.nonempty for cc in range(FIRST_DATA_COL, last_col + 1))
        if lab_b[0] is None and lab_c[0] is None and not has_values:
            continue
        stub.append({"row": r, "group": group, "category": lab_b[0], "subcategory": lab_c[0],
                     "label": lab_c[0] or lab_b[0], "level": "subcategory" if lab_c[0] else "category",
                     "label_from_number_format": lab_b[1] or lab_c[1], "has_values": has_values})

    footnotes = [clean(sh.value(r, 1)) for r in range(source_row, (sh.bbox() or (0, 0, 0, 0))[2] + 1)
                 if sh.value(r, 1) is not None] if source_row else []

    return {
        "table_no": table_no, "title": title, "unit_rows": unit_rows, "unit_row": unit_row,
        "source_row": source_row, "last_col": last_col, "header_top_row": top_row,
        "questions": clean(sh.value(2, 1) or ""), "notes": clean(sh.value(5, 1) or ""),
        "notes_raw": str(sh.value(5, 1) or ""), "questions_raw": str(sh.value(2, 1) or ""),
        "measures": measures, "stub": stub, "footnotes": footnotes,
    }


def cell_state(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "blank"
    if is_number(v):
        return "numeric"
    s = v.strip()
    if s == "..":
        return "dotdot"
    if s in ("…", "..."):
        return "ellipsis"
    if s in CODES:
        return "quality_code"
    try:
        float(s.replace(",", ""))
        return "numeric_text"
    except ValueError:
        return "other_text"


# ---------------------------------------------------
# 4. Declared interpretation (judgement, with evidence)
# ---------------------------------------------------

SHEET_ROLES = {
    "Cover page": "presentation (cover)",
    "Title Page": "presentation (title page)",
    "Notes": "documentation (population, collection, weighting, quality codes)",
    "Index": "lookup (table index)",
}

DIMENSION_CLASS = {
    "All SMEs": "all_smes", "Employment Size": "employment_size", "Region": "geography_region",
    "CMA Level": "geography_cma", "Industries": "industry", "Industry": "industry",
    "Location": "rural_urban", "Export Status": "export_status",
    "International Business Activity": "international_activity", "Growth": "sales_growth",
    "New Markets": "expansion_intention", "Age of Business": "business_age",
    "Innovation Activities developed or introduced": "innovation",
    "New Technology Adoption": "technology_adoption", "E-Commerce": "ecommerce",
    "Intellectual Property": "intellectual_property",
    "Age of Primary Decision Maker": "owner_age",
    "Highest Level of Education Attained by the Primary Decision Maker": "owner_education",
    "Place of birth": "owner_birthplace", "Female Ownership": "female_ownership",
    "Demographics": "ownership_demographics", "Potential Borrowers": "potential_borrowers",
    "Rejected (Debt) Borrowers": "rejected_borrowers", "Procurement": "procurement",
    "Co-operatives": "cooperatives", "Clean Tech": "clean_tech",
    "Social Enterprises": "social_enterprises", "Legal Status": "legal_status",
}
OWNER_DEMOGRAPHIC_CLASSES = {"owner_age", "owner_education", "owner_birthplace",
                             "female_ownership", "ownership_demographics"}
DEFINITIONAL_CLASSES = {"potential_borrowers", "rejected_borrowers"}
REGION_LEVEL = {
    "Atlantic": "multi_province_region", "Quebec": "province", "Ontario": "province",
    "Northern Ontario": "sub_provincial_region", "Southern Ontario": "sub_provincial_region",
    "Prairies": "multi_province_region", "Manitoba": "province", "Saskatchewan": "province",
    "Alberta": "province", "British Columbia and Territories": "province_plus_territories",
}
# Dimensions the project asked to test for, including ones the workbook may not carry.
REQUIRED_DIMENSIONS = ["all_smes", "geography_region", "province", "geography_cma", "industry_naics_group",
                       "industry_special_aggregate", "employment_size", "revenue_band"]

ALL = "All SMEs in the row segment (question applicable to all respondents)"

# Denominators for every block of every Core/Contextual table. Evidence cites the sheet's
# "Notes" cell (A5), the question text (A2) or a reproduced data check.
# multiple: "yes"/"no"/"n/a"; exclusive: whether the block's categories are mutually exclusive.
ANNOTATIONS = {
    ("1", 1): dict(universe=ALL, evidence="Table 1 notes: Q10, Q16, Q21, Q27, Q35, Q39, Q43, Q45 apply to all respondents", multiple="not applicable", exclusive="yes (complement of block 2)"),
    ("1", 2): dict(universe=ALL, evidence="Complement of block 1 (reconciliation R_T01_REQUEST_COMPLEMENT)", multiple="not applicable", exclusive="yes (complement of block 1)"),
    ("1", 3): dict(universe="SMEs in the row segment that did not request external financing in 2023",
                   evidence="FACT: DATA: the six reasons sum to 100 (R_T01_REASONS) and 'Financing not required' (All SMEs 81.4) exceeds the non-request rate (50.7), so the base cannot be all SMEs. The sheet note says Q47 is 'applicable to all respondents', which conflicts with the data.",
                   multiple="no (sums to 100)", exclusive="yes"),
    ("2", 1): dict(universe="SMEs in the row segment approved for external debt financing in 2023", evidence="Table 2 notes (Q34); question allows multiple responses", multiple="yes", exclusive="no"),
    ("3", 1): dict(universe=ALL, evidence="Table 3 notes: Q10, Q16, Q21, Q27 apply to all respondents", multiple="not applicable", exclusive="not applicable"),
    ("3", 2): dict(universe="SMEs in the row segment that requested debt financing in 2023 (outcome of the largest request)", evidence="Table 3 notes: Q11/Q17/Q22/Q28 apply only to requesters; header 'Outcome of largest debt financing request'", multiple="no", exclusive="yes (R_T03_FULL_PARTIAL, R_T03_OUTCOMES)"),
    ("4", 1): dict(universe=ALL, evidence="Table 4 notes: request questions apply to all respondents", multiple="yes (instrument request rates sum above 'Overall')", exclusive="no"),
    ("4", 2): dict(universe="SMEs in the row segment that requested the instrument ('Overall': any debt instrument)", evidence="Table 4 notes: outcome questions apply only to requesters; '* Overall approval rates include both full and partial approvals'", multiple="not applicable", exclusive="not applicable"),
    ("5", 1): dict(universe="SMEs in the row segment that requested debt financing in 2023 (largest request per instrument)", evidence="Table 5 notes; questions ask the dollar amount of the largest request. Totals are weighted sums; averages are per requesting SME; 'authorized-to-requested' is total authorized / total requested x 100", multiple="not applicable", exclusive="not applicable"),
    ("6", 1): dict(universe="SMEs in the row segment approved for a non-residential mortgage in 2023", evidence="Table 6 notes (Q14, Q15); values are average interest rate (%) and average term (months)", multiple="not applicable", exclusive="not applicable"),
    ("6", 2): dict(universe="SMEs in the row segment approved for a new business line of credit in 2023", evidence="Table 6 notes (Q20); value is an average interest rate", multiple="not applicable", exclusive="not applicable"),
    ("6", 3): dict(universe="SMEs in the row segment approved for a term loan in 2023", evidence="Table 6 notes (Q25, Q26)", multiple="not applicable", exclusive="not applicable"),
    ("6", 4): dict(universe="SMEs in the row segment approved for a business credit card or higher limit in 2023", evidence="Table 6 notes (Q31)", multiple="not applicable", exclusive="not applicable"),
    ("6", 5): dict(universe="SMEs in the row segment approved for a mortgage, term loan, line of credit or credit card in 2023", evidence="Table 6 notes (Q33); collateral types allow multiple responses", multiple="yes", exclusive="no"),
    ("7", 1): dict(universe="SMEs in the row segment that requested debt financing in 2023", evidence="Table 7 notes (Q37 as printed; question text shows Q32)", multiple=None, exclusive=None),
    ("8", 1): dict(universe="SMEs in the row segment that requested debt financing and were turned down in 2023", evidence="Table 8 notes (Q13, Q19, Q24, Q30); multiple responses", multiple="yes", exclusive="no"),
    ("9", 1): dict(universe=ALL, evidence="Table 9 notes (Q35 applies to all respondents)", multiple="not applicable", exclusive="not applicable"),
    ("9", 2): dict(universe="SMEs in the row segment that requested lease financing in 2023", evidence="Table 9 notes (Q36)", multiple="no", exclusive="yes (R_T09_FULL_PARTIAL)"),
    ("10", 1): dict(universe="SMEs in the row segment that requested lease financing in 2023", evidence="Table 10 notes (Q35, Q36)", multiple="not applicable", exclusive="not applicable"),
    ("11", 1): dict(universe="SMEs in the row segment that requested lease financing in 2023", evidence="Table 11 notes (Q38)", multiple=None, exclusive=None),
    ("12", 1): dict(universe=ALL, evidence="Table 12 notes (Q39 applies to all respondents)", multiple="not applicable", exclusive="not applicable"),
    ("12", 2): dict(universe="Approval rate: SMEs that requested trade credit (Q40); amounts: SMEs approved for trade credit (Q42)", evidence="Table 12 notes", multiple="not applicable", exclusive="not applicable"),
    ("13", 1): dict(universe=ALL, evidence="Table 13 notes (Q43 applies to all respondents)", multiple="not applicable", exclusive="not applicable"),
    ("13", 2): dict(universe="SMEs in the row segment that requested equity financing in 2023", evidence="Table 13 notes", multiple="not applicable", exclusive="not applicable"),
    ("13", 3): dict(universe="SMEs in the row segment that requested equity financing in 2023", evidence="Table 13 notes", multiple="not applicable", exclusive="not applicable"),
    ("14", 1): dict(universe=ALL, evidence="Table 14 notes (Q45 applies to all respondents)", multiple="not applicable", exclusive="not applicable"),
    ("14", 2): dict(universe="SMEs that requested government financing in 2023 (the note also says 'or approved ... in 2020', recorded as printed)", evidence="Table 14 notes (Q46); the 2020 reference is probably carried over from an earlier cycle", multiple="not applicable", exclusive="not applicable"),
    ("15", 1): dict(universe="SMEs whose debt request of that type was fully authorized in 2023", evidence="Table 15 notes (Q12, Q18, Q23, Q29)", multiple="no", exclusive="yes (amount classes)"),
    ("15", 2): dict(universe="SMEs whose lease request was fully authorized in 2023", evidence="Table 15 notes (Q37)", multiple="no", exclusive="yes (amount classes)"),
    ("15", 3): dict(universe="SMEs provided trade credit in 2023", evidence="Table 15 notes (Q41)", multiple="no", exclusive="yes (amount classes)"),
    ("15", 4): dict(universe="SMEs that requested new equity financing in 2023", evidence="Table 15 notes (Q44)", multiple="no", exclusive="yes (amount classes)"),
    ("16", 1): dict(universe=ALL, evidence="Table 16 notes (Q5); age classes plus 'Average age' (years)", multiple="no", exclusive="yes (R_T16_ESTABLISHED)"),
    ("16", 2): dict(universe=ALL, evidence="Table 16 notes (Q6); classes plus average years selling", multiple="no", exclusive="yes (R_T16_SELLING)"),
    ("16", 3): dict(universe=ALL, evidence="Table 16 notes (Q7); the three methods do not sum to 100 (other methods not shown)", multiple="no", exclusive="yes (not exhaustive)"),
    ("17", 1): dict(universe="SMEs in the row segment that started their business from scratch", evidence="Table 17 notes (Q8); multiple responses", multiple="yes", exclusive="no"),
    ("18", 1): dict(universe="SMEs in the row segment that bought or acquired their business", evidence="Table 18 notes (Q9); multiple responses", multiple="yes", exclusive="no"),
    ("19", 1): dict(universe="Past growth: SMEs that sold goods or services in the last three years (Q48); expected growth: all respondents (Q49)", evidence="Table 19 notes", multiple="no", exclusive="yes (R_T19_PAST, R_T19_EXPECTED)",
                    measure_rules=[("over the last three years", "SMEs in the row segment that sold goods or services in the last three years (Q48)"),
                                   ("over the next three years", ALL)]),
    ("20", 1): dict(universe=ALL, evidence="Table 20 notes (Q50); each obstacle is rated on four levels", multiple="no (within an obstacle)", exclusive="yes within an obstacle (R_T20_SEVERITY)"),
    ("26", 1): dict(universe=ALL, evidence="Table 26 notes (Q62)", multiple="not applicable", exclusive="not applicable"),
    ("26", 2): dict(universe="Markets: per note Q62 applies to all respondents; whether the base is all SMEs or SMEs intending to expand is not stated (UNRESOLVED). Export destinations: SMEs intending to expand outside Canada (Q63)", evidence="Table 26 notes; multiple responses", multiple="yes", exclusive="no",
                    measure_rules=[("outside of canada (export sales)", "SMEs in the row segment intending to expand sales outside Canada (Q63)")]),
    ("38", 1): dict(universe=ALL, evidence="Table 38 notes (Q93)", multiple="no", exclusive="yes (R_T38_INTENT)"),
    ("38", 2): dict(universe=ALL, evidence="Table 38 notes (Q93)", multiple="no", exclusive="yes (R_T38_INTENT)"),
    ("38", 3): dict(universe="SMEs in the row segment intending to sell, transfer or close within five years (Q94)", evidence="Table 38 notes; the six intentions sum to 100 (R_T38_INTENTION_TYPE). Row-14 labels 'Latin America', 'Europe', 'China' and 'Other' above these columns are stray text (see anomalies)", multiple="no", exclusive="yes"),
    ("38", 4): dict(universe=ALL, evidence="Table 38 notes (Q93)", multiple="no", exclusive="yes (R_T38_INTENT)"),
    ("49", 1): dict(universe="Survey respondents in the row segment (unweighted count; sample support, not an estimate)", evidence="Table 49 header 'Number of Respondents'; All SMEs = 11,210 = responses in the Notes sheet", multiple="not applicable", exclusive="not applicable"),
    ("49", 2): dict(universe="Estimated number of enterprises in the survey's target population, in the row segment", evidence="Table 49 title 'Weighted Enterprise Counts'; Notes: sampling weights represent the target population", multiple="not applicable", exclusive="not applicable"),
    ("49", 3): dict(universe="Share of the All-SMEs weighted count for 126 of 132 rows. Six special-population rows use another base: ICT, co-operatives, clean technology and federal-procurement SMEs print 100; charity and not-for-profit split the non-profit social-enterprise group (59.7 + 40.3)",
                    evidence="FACT: DATA: R_T49_SHARE (126 rows match 100 x weighted / All-SMEs weighted; 6 do not). The base of the six is not documented", multiple="not applicable", exclusive="not applicable"),
    ("50", 1): dict(universe="Estimated number of enterprises (weighted), cross-tabulated by employment size", evidence="Table 50 title", multiple="not applicable", exclusive="not applicable"),
    ("50", 2): dict(universe="Estimated number of enterprises (weighted), cross-tabulated by region (industry rows only; region-by-region cells are blank)", evidence="Table 50 title and layout", multiple="not applicable", exclusive="not applicable"),
    ("A", 1): dict(universe="Incorporated SMEs in the row segment with complete tax data (89.8% of all SMEs); CRA administrative data", evidence="Table A notes and footnote", multiple="not applicable", exclusive="not applicable"),
}

# Relevance to the SME acquisition project (DECISION). Core/Contextual rows carry the
# supporting business question, decision, usable grain, concept, comparability, cautions
# and reference-year alignment.
BQ_FOCUS = "Where should SME acquisition effort be focused?"
BQ_ROUTE = "How should SME leads be routed to relationship-manager specializations?"
GRAIN_STD = ("One-way segments only: Canada; 5 regions (Quebec, Ontario, Manitoba, Saskatchewan, Alberta also individually; "
             "Ontario north/south); 24 CMAs by name; 10 NAICS-based industry groups plus overlapping special aggregates; "
             "4 employment-size bands; other business characteristics. No cross-tabulations.")
CAUTION_STD = ("Survey estimate with sampling error: use A-D, flag E, exclude F; '..' is unavailable. Quality A-D means each estimate "
               "is individually reliable, not that two segments can be ranked: differences smaller than about twice their combined "
               "standard error (SE/CV printed per estimate) should not order segments. Percentages are prevalence within the row "
               "segment's applicable population, not counts of businesses. One-way segments cannot be combined into CMA x industry "
               "x size cells.")
ALIGN_STD = "Reference year 2023 (collection Feb-Jun 2024). Suitable for structural context; about two years older than 1097 (Dec 2025)."

RELEVANCE = {
    "49": ("Core", "market size", "Estimated number of in-scope SMEs by segment (size, region, CMA, industry and others): the survey's own market-scale measure", BQ_FOCUS, "Size and rank segments by the number of in-scope SMEs"),
    "50": ("Core", "market size", "Estimated SMEs cross-tabulated size x region, size x industry and industry x region: the only two-way cells in the workbook", BQ_FOCUS, "Size segment combinations (e.g. industry within region)"),
    "1": ("Core", "financing behaviour (demand)", "Share of SMEs that requested any external financing, and why others did not", BQ_FOCUS, "Prioritize segments with higher financing demand"),
    "3": ("Core", "financing behaviour (debt demand and approval)", "Debt request rate and outcome of the largest request", BQ_FOCUS, "Prioritize segments with debt demand; understand approval environment"),
    "4": ("Core", "financing behaviour (product demand)", "Request and approval rates by debt instrument (mortgage, line of credit, term loan, credit card)", BQ_ROUTE, "Match product specializations to segment demand"),
    "5": ("Core", "financing behaviour (deal size)", "Total and average amounts requested and authorized by instrument", BQ_FOCUS, "Estimate typical deal size by segment"),
    "2": ("Contextual", "competitive landscape", "Providers of approved debt financing (chartered banks, credit unions, government, online lenders)", BQ_FOCUS, "Understand provider mix by segment"),
    "6": ("Contextual", "financing terms", "Average interest rates, terms and collateral requirements", BQ_ROUTE, "Context for product pricing and structuring"),
    "7": ("Contextual", "financing needs", "Intended uses of debt financing", BQ_ROUTE, "Context for product specialization"),
    "8": ("Contextual", "financing outcomes", "Reasons lenders gave for turning down requests", BQ_FOCUS, "Context for credit-risk profile of segments"),
    "9": ("Contextual", "financing behaviour (leasing)", "Lease request and approval", BQ_ROUTE, "Context for leasing specialization"),
    "10": ("Contextual", "financing behaviour (leasing)", "Lease amounts requested and authorized", BQ_ROUTE, "Context for leasing deal size"),
    "11": ("Contextual", "financing needs (leasing)", "Assets acquired through leasing", BQ_ROUTE, "Context for leasing specialization"),
    "12": ("Contextual", "financing behaviour (trade credit)", "Trade credit use, approval and amounts", BQ_FOCUS, "Context for working-capital needs"),
    "13": ("Contextual", "financing behaviour (equity)", "Equity financing requests and amounts", BQ_FOCUS, "Context: equity is rare (1.4% of SMEs requested)"),
    "14": ("Contextual", "financing behaviour (government)", "Government financing requests, approval and amounts", BQ_FOCUS, "Context for government-backed programs"),
    "15": ("Contextual", "financing behaviour (deal size distribution)", "Approved amount classes by financing type", BQ_FOCUS, "Context for deal-size distribution"),
    "16": ("Contextual", "business characteristics", "Business age and how the business was acquired", BQ_FOCUS, "Context for life-stage targeting"),
    "17": ("Contextual", "financing behaviour (start-up)", "Sources of start-up financing", BQ_FOCUS, "Context for new-business segment"),
    "18": ("Contextual", "financing behaviour (acquisition)", "Sources of financing to buy a business", BQ_FOCUS, "Context for acquisition-financing need"),
    "19": ("Contextual", "growth expectations", "Past and expected sales growth distribution", BQ_FOCUS, "Context for growth-oriented segments"),
    "20": ("Contextual", "business needs", "Obstacles to growth rated by severity (including access to financing)", BQ_FOCUS, "Context for segment needs"),
    "26": ("Contextual", "growth expectations", "Intentions to expand into new markets", BQ_FOCUS, "Context for expansion-financing need"),
    "38": ("Contextual", "business characteristics (succession)", "Intentions to sell, transfer or close within five years", BQ_FOCUS, "Context for ownership-transition financing"),
    "A": ("Contextual", "business characteristics (financial scale)", "Average and median revenue, expenses, profit, assets and liabilities (CRA data)", BQ_FOCUS, "Context for revenue scale of segments"),
}
OUT_OF_SCOPE_REASON = {
    **{n: "Export and sales-destination behaviour; not an acquisition or routing input" for n in ("21", "22", "23", "24", "25", "27")},
    **{n: "Government procurement; not an acquisition or routing input" for n in ("28", "29")},
    **{n: "Technology adoption, online sales and consultants; not a financing or market-scale measure" for n in ("30", "31", "32", "33", "34")},
    **{n: "Owner demographics; must not drive targeting or routing" for n in ("35", "36", "37")},
    **{n: "Innovation and intellectual property; not an acquisition or routing input" for n in ("39", "40", "41")},
    **{n: "Social-enterprise sub-population (largely non-profit); outside the business-banking scope" for n in ("42", "43", "44", "45", "46", "47", "48")},
}

# '..' contexts that are not applicable by definition (INFERENCE from the table's population).
DOTDOT_DEFINITIONAL_RULES = [
    (None, "Potential Borrowers", None, None, "Potential borrowers did not request financing, so request-conditional measures do not apply"),
    ("1", "Rejected (Debt) Borrowers", None, {2, 3}, "Rejected borrowers requested financing, so 'did not request' measures do not apply"),
    ("A", "Legal Status", ("Unincorporated", "Sole Proprietorship"), None, "Table A covers incorporated SMEs only (footnote)"),
    ("24", "Export Status", ("Exporter",), None, "Reasons for not exporting apply only to non-exporters"),
    ("25", "Export Status", ("Non-exporter",), None, "Export obstacles apply only to exporters"),
    ("22", "Export Status", ("Non-exporter",), None, "Export-sales measures apply only to exporters"),
    ("23", "Export Status", ("Non-exporter",), None, "Export type and destination apply only to exporters"),
    ("26", "New Markets", ("Does not intend to expand sales into new markets",), None, "Target markets apply only to SMEs intending to expand"),
    ("32", "New Technology Adoption", ("Adopted at least one new technology over the last three years",), None, "Reasons for not investing apply only to non-adopters"),
    ("19", "Age of Business", ("2021-2023 (2 years or younger)",), None, "Three-year past growth does not apply to businesses two years old or younger"),
    (("45", "47"), "Social Enterprises", ("Not-for-profit", "Charity"), None, "Question applies to for-profit social enterprises only"),
    (("42", "43", "44", "46", "48"), "Social Enterprises", ("For-profit",), None, "Question applies to non-profit social enterprises only"),
]


def dotdot_context(table_no, group, label, block_no):
    for tables, grp, labels, blocks, why in DOTDOT_DEFINITIONAL_RULES:
        if tables is not None and table_no not in (tables if isinstance(tables, tuple) else (tables,)):
            continue
        if group != grp or (labels is not None and label not in labels) or (blocks is not None and block_no not in blocks):
            continue
        return "not_applicable_by_definition", why
    return "unavailable_reason_undocumented", "Not explained by the table's population; consistent with small-sample suppression, which the source does not document"


def classify_measure(table_no, m):
    path = " > ".join(m["header_path"]).lower()
    tok = m["unit_token"]
    if table_no == "49":
        return {1: ("unweighted_respondent_count", "respondents"),
                2: ("weighted_count", "enterprises (weighted estimate)"),
                3: ("percentage_share_of_total", "percent (0-100)")}[m["block_no"]]
    if table_no == "50":
        return "weighted_count", "enterprises (weighted estimate)"
    if table_no == "A":
        return ("average_monetary" if tok == "Average" else "median_monetary"), "Canadian dollars"
    if tok == "$":
        return ("average_monetary" if "average" in path else "weighted_total_monetary"), "Canadian dollars"
    if tok == "Months":
        return "average_duration", "months"
    if tok in ("Average", "average"):
        return "average_duration", "years (inferred from header text)"
    if "interest rate" in path:
        return "average_rate", "percent (interest rate)"
    if "authorized-to-requested" in path:
        return "ratio_of_totals", "percent (0-100)"
    return "percentage_of_universe", "percent (0-100)"


def naics_codes(label):
    return ";".join(re.findall(r"NAICS\s+([\d\-]+)", label or ""))


def norm_place(s):
    s = unicodedata.normalize("NFKD", str(s).replace(NBSP, " ")).encode("ascii", "ignore").decode()
    s = s.split(",")[0]
    return re.sub(r"\s*-\s*", "-", s).strip().lower()


NBSP = " "

# Stray header text found in the published workbook (kept as-is, documented).
HEADER_ANOMALIES = {
    ("38", "M14"): "Latin America", ("38", "P14"): "Europe", ("38", "S14"): "China", ("38", "Y14"): "Other",
}


def bisect_code(err, bounds):
    """Quality code implied by an error value under a threshold scheme (5 upper bounds for A-E)."""
    for code, upper in zip(CODES, bounds):
        if err < upper:
            return code
    return "F"


def years_and_periods(text):
    years = sorted(set(re.findall(r"\b(20\d\d|19\d\d)\b", text)))
    rel = sorted(set(m.lower() for m in re.findall(
        r"(last (?:3|three) years|next (?:3|three) years|next five years|as of december 2023|last three years)", text, re.I)))
    return ";".join(years), ";".join(rel)


# ---------------------------------------------------
# Main
# ---------------------------------------------------

def main():
    start = time.time()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    banner("ISED SFGSME 2023 WORKBOOK — PHASE C VALIDATION")

    # ---------------------------------------------------
    # 1. Source integrity (before)
    # ---------------------------------------------------
    banner("1. SOURCE INTEGRITY (BEFORE)")
    manifest = pd.read_csv(MANIFEST_FILE, dtype=str).fillna("")
    rel = RAW_FILE.relative_to(ROOT).as_posix()
    mrow = manifest[manifest.relative_path == rel]
    sha_before, stat_before = sha256_file(RAW_FILE), RAW_FILE.stat()
    gate("workbook_listed_once_in_manifest", len(mrow) == 1)
    m0 = mrow.iloc[0]
    gate("workbook_sha256_matches_manifest_before", m0.sha256 == sha_before, sha_before[:16] + "...")
    gate("workbook_size_matches_open_government_portal", stat_before.st_size == PORTAL_FILE_BYTES, f"{stat_before.st_size:,} bytes")
    gate("manifest_provenance_recorded_for_workbook",
         all(m0[f] for f in ("table_title", "table_page_url", "metadata_reference_period"))
         and "to be recorded" not in m0.provenance_note.lower())

    # ---------------------------------------------------
    # 2. Package + workbook structure
    # ---------------------------------------------------
    banner("2. PACKAGE AND WORKBOOK STRUCTURE")
    pkg = inspect_package()
    wb_f = openpyxl.load_workbook(RAW_FILE, data_only=False)
    wb_v = openpyxl.load_workbook(RAW_FILE, data_only=True)
    sheets_f = [Sheet(ws) for ws in wb_f.worksheets]
    sheets_v = [Sheet(ws) for ws in wb_v.worksheets]
    n_sheets = len(sheets_f)
    gate("sheet_count_consistent_xml_app_openpyxl",
         len(pkg["sheets"]) == pkg["app_titles_of_parts"] == n_sheets, f"{n_sheets} worksheets")
    gate("sheet_names_and_order_match_xml", [s.title for s in sheets_f] == [s["name"] for s in pkg["sheets"]])
    states = Counter(ws.sheet_state for ws in wb_f.worksheets)
    gate("all_sheets_inventoried_including_hidden", sum(states.values()) == n_sheets, dict(states))
    gate("no_external_links_pivots_charts_tables_macros",
         pkg["external_link_parts"] == pkg["pivot_parts"] == pkg["chart_parts"] == pkg["table_parts"] == pkg["vba_parts"] == 0)
    user_names = [d for d in pkg["defined_names"] if d["kind"] == "user_named_range"]
    gate("no_user_named_ranges", len(user_names) == 0, f"{len(pkg['defined_names'])} defined names, all housekeeping")

    formulas = sum(1 for sh in sheets_f for c in sh.nonempty.values()
                   if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")))
    mismatch = 0
    for a, b in zip(sheets_f, sheets_v):
        va = {k: c.value for k, c in a.nonempty.items()}
        vb = {k: c.value for k, c in b.nonempty.items()}
        mismatch += len(set(va.items()) ^ set(vb.items()))
    gate("no_formulas", formulas == 0)
    gate("formula_and_cached_value_reads_identical", mismatch == 0)

    # Index sheet: table number -> title.
    idx = next(s for s in sheets_f if s.title == "Index")
    index_titles = {}
    for (r, c), cell in idx.nonempty.items():
        if c == 2 and r > 3:
            index_titles[str(cell.value).strip()] = clean(idx.value(r, 3) or "")

    # Notes sheet facts (text).
    notes = next(s for s in sheets_f if s.title == "Notes")
    notes_text = " ".join(clean(c.value) for c in notes.nonempty.values() if isinstance(c.value, str))
    m_sample = re.search(r"sample size for this survey was ([\d,]+)", notes_text)
    m_resp = re.search(r"responses obtained from ([\d,]+) SMEs", notes_text)
    m_rate = re.search(r"response rate, as a percentage of in-scope businesses, is ([\d.]+)%", notes_text)
    notes_sample = int(m_sample.group(1).replace(",", "")) if m_sample else None
    notes_responses = int(m_resp.group(1).replace(",", "")) if m_resp else None
    notes_rate = float(m_rate.group(1)) if m_rate else None
    cover = next(s for s in sheets_f if s.title == "Cover page")
    cover_text = " ".join(clean(c.value) for c in cover.nonempty.values() if isinstance(c.value, str))
    release = re.search(r"Released on ([A-Z][a-z]+ \d{1,2}, \d{4})", cover_text)

    # ---------------------------------------------------
    # 3. Parse every table sheet
    # ---------------------------------------------------
    banner("3. TABLE DISCOVERY")
    tables, measures, tidy = [], [], []
    xml_dims = {s["name"]: s["xml_dimension"] for s in pkg["sheets"]}
    wb_rows = []
    for order, sh in enumerate(sheets_f, start=1):
        ws = sh.ws
        bb = sh.bbox()
        xml_dim = xml_dims[sh.title]
        min_c, min_r, max_c, max_r = openpyxl.utils.cell.range_boundaries(xml_dim)
        within = bb is None or (min_r <= bb[0] and min_c <= bb[1] and bb[2] <= max_r and bb[3] <= max_c)
        is_table = bool(re.match(r"Table \w+$", sh.title))
        t = parse_table(sh) if is_table else None
        text_blob = " ".join(clean(c.value) for c in sh.nonempty.values() if isinstance(c.value, str))
        yrs, rels = years_and_periods(text_blob)
        numeric_cells = sum(is_number(c.value) for c in sh.nonempty.values())
        wb_rows.append({
            "sheet_order": order, "sheet_name": sh.title, "sheet_state": ws.sheet_state,
            "xml_part": next(s["xml_part"] for s in pkg["sheets"] if s["name"] == sh.title),
            "xml_dimension": xml_dim, "used_range_openpyxl": sh.used_range,
            "nonempty_bounding_box": (f"{col_letter(bb[1])}{bb[0]}:{col_letter(bb[3])}{bb[2]}" if bb else None),
            "nonempty_within_xml_dimension": within,
            "nonempty_cells": len(sh.nonempty), "numeric_cells": numeric_cells,
            "text_cells": sum(isinstance(c.value, str) for c in sh.nonempty.values()),
            "nonempty_rows": len({r for r, _ in sh.nonempty}), "nonempty_cols": len({c for _, c in sh.nonempty}),
            "merged_ranges": len(sh.merged),
            "hidden_rows": sum(1 for d in ws.row_dimensions.values() if d.hidden),
            "hidden_cols": sum(1 for d in ws.column_dimensions.values() if d.hidden),
            "freeze_panes": ws.freeze_panes, "formulas": 0,
            "defined_names_local": sum(1 for d in pkg["defined_names"] if d["local_sheet_id"] == str(order - 1)),
            "autofilter_ref": ws.auto_filter.ref,
            "tables": len(ws.tables), "charts": len(ws._charts), "images": len(ws._images),
            "conditional_format_rules": sum(len(cf.rules) for cf in ws.conditional_formatting),
            "data_validations": len(ws.data_validations.dataValidation),
            "comments": sum(1 for c in sh.cells.values() if c.comment is not None),
            "sheet_role": "data table" if is_table else SHEET_ROLES.get(sh.title, "unknown"),
            "apparent_title": (f"Table {t['table_no']}: {t['title']}" if t else
                               {"Cover page": clean(" ".join(str(sh.value(r, 1)) for r in (13, 14, 15))),
                                "Title Page": clean(sh.value(7, 1)), "Notes": clean(sh.value(1, 1)),
                                "Index": "Index of tables"}.get(sh.title)),
            "apparent_subtitle": (t["questions"][:200] if t else
                                  {"Cover page": clean(sh.value(18, 1)), "Title Page": clean(sh.value(9, 1)),
                                   "Notes": "Target population, data collection, weighting, data accuracy"}.get(sh.title)),
            "years_mentioned": yrs, "relative_periods_mentioned": rels,
            "apparent_units": (";".join(sorted({m["unit_token"] for m in t["measures"]})) if t else None),
            "content_type": ("data" if is_table else "presentation" if "presentation" in SHEET_ROLES.get(sh.title, "")
                             else "lookup" if sh.title == "Index" else "documentation"),
        })
        if not t:
            continue

        table_no = t["table_no"]
        tid = f"ISED_T{int(table_no):02d}" if table_no.isdigit() else f"ISED_T{table_no}"
        stub_keys = [(s["group"], s["label"]) for s in t["stub"]]
        roles = {}
        for mi, m in enumerate(t["measures"], start=1):
            for role, col in (("value", m["value_col"]), ("quality", m["quality_col"]), ("error", m["error_col"])):
                if col is not None:
                    roles.setdefault(col, []).append((mi, role))
        data_cols = {c for (r, c) in sh.nonempty if t["stub"] and t["stub"][0]["row"] <= r <= t["stub"][-1]["row"] and c >= FIRST_DATA_COL}
        merged_in_data = sum(1 for mr in sh.merged if mr.max_row >= t["stub"][0]["row"] and mr.min_row <= t["stub"][-1]["row"]
                             and mr.max_col >= FIRST_DATA_COL)
        tables.append({"table_no": table_no, "table_id": tid, "sheet": sh.title, "sheet_order": order, "t": t, "sh": sh,
                       "stub_keys": stub_keys, "col_roles": roles, "data_cols": data_cols, "merged_in_data": merged_in_data})

        for mi, m in enumerate(t["measures"], start=1):
            est_type, unit = classify_measure(table_no, m)
            mid = f"{tid}_M{mi:03d}"
            bid = f"{tid}_B{m['block_no']:02d}"
            measures.append({"table_no": table_no, "table_id": tid, "measure_idx": mi, "measure_id": mid,
                             "block_no": m["block_no"], "block_id": bid, "block_title": m["block_title"],
                             "value_col": col_letter(m["value_col"]),
                             "quality_col": col_letter(m["quality_col"]) if m["quality_col"] else None,
                             "error_col": col_letter(m["error_col"]) if m["error_col"] else None,
                             "error_type": m["error_type"], "unit_token": m["unit_token"],
                             "header_path": " > ".join(m["header_path"]), "estimate_type": est_type, "unit": unit})
            for s in t["stub"]:
                vc = sh.cells.get((s["row"], m["value_col"]))
                v = vc.value if vc is not None else None
                q = sh.value(s["row"], m["quality_col"]) if m["quality_col"] else None
                e = sh.value(s["row"], m["error_col"]) if m["error_col"] else None
                tidy.append({"table_no": table_no, "table_id": tid, "measure_idx": mi, "measure_id": mid,
                             "block_no": m["block_no"], "unit_token": m["unit_token"], "estimate_type": est_type,
                             "stub_row": s["row"], "group": s["group"], "label": s["label"], "level": s["level"],
                             "value": float(v) if is_number(v) else None, "value_state": cell_state(v),
                             "value_number_format": vc.number_format if vc is not None else None,
                             "quality": q.strip() if isinstance(q, str) else q, "quality_state": cell_state(q),
                             "error": float(e) if is_number(e) else None, "error_state": cell_state(e),
                             "error_type": m["error_type"]})

    cells = pd.DataFrame(tidy)
    meas = pd.DataFrame(measures)
    wbinv = pd.DataFrame(wb_rows)
    tables_by_no = {tb["table_no"]: tb for tb in tables}
    print(f"  tables: {len(tables)} | measures: {len(meas)} | blocks: {meas.block_id.nunique()} | cells: {len(cells):,}")

    gate("every_table_sheet_has_title_unit_row_and_footnotes",
         all(tb["t"]["table_no"] and tb["t"]["unit_rows"] == [16] and tb["t"]["source_row"] for tb in tables))
    gate("sheet_name_matches_title_number", all(tb["sheet"] == f"Table {tb['table_no']}" for tb in tables))
    gate("index_lists_every_table_and_only_those", set(index_titles) == set(tables_by_no), f"{len(index_titles)} index entries")
    gate("every_table_and_measure_has_stable_unique_id",
         len({tb["table_id"] for tb in tables}) == len(tables) and meas.measure_id.is_unique)
    ref_stub = tables_by_no["1"]["stub_keys"]
    same_stub = [tb["table_no"] for tb in tables if tb["table_no"] != "50" and tb["stub_keys"] == ref_stub]
    gate("row_stub_identical_across_tables_1_to_49_and_A", len(same_stub) == len(tables) - 1, f"{len(same_stub)} tables")
    gate("stub_labels_unique_within_each_table", all(len(set(tb["stub_keys"])) == len(tb["stub_keys"]) for tb in tables))
    gate("every_data_column_belongs_to_exactly_one_measure_role",
         all(all(len(v) == 1 for v in tb["col_roles"].values()) and tb["data_cols"] <= set(tb["col_roles"]) for tb in tables))
    gate("measure_column_ranges_do_not_overlap",
         all(len(tb["col_roles"]) == sum(len(v) for v in tb["col_roles"].values()) for tb in tables))
    gate("every_value_column_has_header_path", (meas.header_path.str.len() > 0).all())
    gate("no_merged_range_intersects_data_area", sum(tb["merged_in_data"] for tb in tables) == 0)
    gate("every_measure_in_tables_1_to_50_has_quality_and_error_columns",
         meas[meas.table_no != "A"][["quality_col", "error_col"]].notna().all().all())
    gate("all_sheet_used_ranges_contain_nonempty_cells", wbinv.nonempty_within_xml_dimension.all())

    # Scale / format gates.
    pct = cells[(cells.unit_token == "%") & (cells.value_state == "numeric")]
    gate("percent_values_stored_on_0_to_100_scale", pct.value.between(0, 100).all() and pct.value.max() > 1,
         f"min {pct.value.min():g}, max {pct.value.max():g}")
    fmts = cells.dropna(subset=["value_number_format"]).groupby(["unit_token", "value_number_format"]).size()
    gate("no_percent_number_format_on_values", not any("%" in f for (_, f) in fmts.index))
    gate("no_display_scaling_in_number_formats", not any(re.search(r"[0#],+(?:[;\"\s)]|$)", f) for (_, f) in fmts.index))
    gate("no_numbers_stored_as_text", (cells.value_state != "numeric_text").all() and (cells.error_state != "numeric_text").all())

    # ---------------------------------------------------
    # 4. Row dimensions
    # ---------------------------------------------------
    banner("4. ROW DIMENSIONS")
    cat_rows = []
    for tno in ("1", "50"):
        for s in tables_by_no[tno]["t"]["stub"]:
            cls = DIMENSION_CLASS.get(s["group"], "unclassified")
            geo_level = (REGION_LEVEL.get(s["label"]) if cls == "geography_region" else "cma" if cls == "geography_cma" else None)
            if cls == "industry":
                cls = "industry_naics_group" if "(NAICS" in s["label"] else "industry_special_aggregate"
            cat_rows.append({"stub_set": "tables_1_49_A" if tno == "1" else "table_50", "stub_row": s["row"],
                             "group": s["group"], "category": s["category"], "subcategory": s["subcategory"],
                             "label": s["label"], "level": s["level"], "dimension_class": cls,
                             "geography_level": geo_level, "naics_codes": naics_codes(s["label"]),
                             "owner_demographic": cls in OWNER_DEMOGRAPHIC_CLASSES,
                             "definitional_segment": cls in DEFINITIONAL_CLASSES,
                             "label_from_number_format": s["label_from_number_format"]})
    catalogue = pd.DataFrame(cat_rows)
    write_csv(catalogue, "row_dimension_catalogue.csv")
    gate("every_stub_row_classified", (catalogue.dimension_class != "unclassified").all())
    main_cat = catalogue[catalogue.stub_set == "tables_1_49_A"]
    dim_counts = main_cat.dimension_class.value_counts().to_dict()
    geo_counts = main_cat.geography_level.value_counts().to_dict()
    print(f"  {dim_counts}")

    cls_map = {(r.group, r.label): (r.dimension_class, r.geography_level) for r in catalogue.itertuples()}
    cells["dimension_class"] = [cls_map[(g, l)][0] for g, l in zip(cells.group, cells.label)]
    cells["geography_level"] = [cls_map[(g, l)][1] for g, l in zip(cells.group, cells.label)]

    # ---------------------------------------------------
    # 5. Quality markers
    # ---------------------------------------------------
    banner("5. QUALITY AND SUPPRESSION MARKERS")
    has_q = cells[cells.table_no != "A"]
    numeric = has_q[has_q.value_state == "numeric"]
    nonnum = has_q[has_q.value_state != "numeric"]
    gate("every_numeric_estimate_has_quality_code_and_error",
         (numeric.quality_state == "quality_code").all() and (numeric.error_state == "numeric").all())
    gate("every_unavailable_or_blank_value_has_no_quality_or_error",
         (nonnum.quality_state == "blank").all() and (nonnum.error_state == "blank").all())
    gate("only_known_value_states_present", set(cells.value_state) <= {"numeric", "dotdot", "blank"},
         {k: int(v) for k, v in cells.value_state.value_counts().items()})

    agree = {}
    for scheme, b in THRESHOLDS.items():
        implied = [bisect_code(e, b[t]) for e, t in zip(numeric.error, numeric.error_type)]
        agree[scheme] = int(sum(a == q for a, q in zip(implied, numeric.quality)))
    n_coded = len(numeric)
    gate("quality_codes_follow_workbook_notes_thresholds", agree["workbook_notes"] == n_coded,
         f"{agree['workbook_notes']:,} of {n_coded:,} agree; IMDB thresholds: {agree['imdb_record']:,}")

    err_ranges = numeric.groupby(["error_type", "quality"]).error.agg(["count", "min", "max"]).reset_index()
    dot = cells[cells.value_state == "dotdot"].copy()
    ctx = [dotdot_context(r.table_no, r.group, r.label, r.block_no) for r in dot.itertuples()]
    dot["context"] = [c[0] for c in ctx]
    dot["rationale"] = [c[1] for c in ctx]
    dot_ctx = (dot.groupby(["table_id", "context", "group", "label", "rationale"]).size().reset_index(name="cells")
               .sort_values(["table_id", "context", "group", "label"]))
    write_csv(dot_ctx, "dotdot_contexts.csv")
    dot_by_ctx = dot.context.value_counts().to_dict()
    cells["dotdot_context"] = None
    cells.loc[dot.index, "dotdot_context"] = dot.context

    blanks = cells[cells.value_state == "blank"]
    structural_blank = blanks[(blanks.table_no == "50") & (blanks.group == "Region") & (blanks.block_no == 2)]
    gate("all_blank_values_are_table_50_region_by_region_block", len(structural_blank) == len(blanks), f"{len(blanks)} blank cells")

    zeros = numeric[numeric.value == 0]
    zeros_def = zeros[zeros.dimension_class.isin(DEFINITIONAL_CLASSES)]
    hundreds_def = numeric[(numeric.value == 100) & numeric.dimension_class.isin(DEFINITIONAL_CLASSES) & (numeric.table_no == "1")]
    all_text = [c.value for sh in sheets_f for c in sh.nonempty.values() if isinstance(c.value, str)]
    ellipsis_cells = sum(1 for v in all_text if v.strip() in ("…", "..."))
    ellipsis_footnotes = sum(1 for v in all_text if v.strip().startswith("…"))
    # Conditional formatting: which conditions does it test, and do they ever fire?
    cf_rules = [(sh.title, rule.type, rule.operator, tuple(rule.formula or ()))
                for sh in sheets_f for cf in sh.ws.conditional_formatting for rule in cf.rules]
    cf_x_rules = [r for r in cf_rules if r[1] == "cellIs" and r[2] == "equal" and r[3] == ('"x"',)]
    x_cells = sum(1 for v in all_text if v.strip().lower() == "x")
    gate("conditional_formatting_only_highlights_x_and_no_x_present", len(cf_x_rules) == len(cf_rules) and x_cells == 0,
         f"{len(cf_rules)} rules, all 'cell equals \"x\"'; {x_cells} cells contain x")
    f_printed = int((numeric.quality == "F").sum())
    e_printed = int((numeric.quality == "E").sum())

    def rng(code, et):
        r = err_ranges[(err_ranges.quality == code) & (err_ranges.error_type == et)]
        return (f"{r['min'].iloc[0]:.4f}-{r['max'].iloc[0]:.4f}" if len(r) else "none")

    wb_meaning = {"A": "Excellent", "B": "Very good", "C": "Good", "D": "Acceptable", "E": "Use with caution", "F": "Not reliable, do not use"}
    imdb_meaning = {"A": "Excellent", "B": "Good", "C": "Average", "D": "Mediocre", "E": "Poor, use with caution", "F": "Not reliable, do not use"}
    qrows = []
    for i, code in enumerate(CODES):
        lo_cv = 0 if i == 0 else THRESHOLDS["workbook_notes"]["CV"][i - 1]
        hi_cv = THRESHOLDS["workbook_notes"]["CV"][i] if i < 5 else None
        lo_se = 0 if i == 0 else THRESHOLDS["workbook_notes"]["SE"][i - 1]
        hi_se = THRESHOLDS["workbook_notes"]["SE"][i] if i < 5 else None
        qrows.append({"marker": code, "marker_type": "quality code (with printed CV or SE)",
                      "official_meaning": f"{wb_meaning[code]} (workbook Notes); '{imdb_meaning[code]}' (IMDB record 2941)",
                      "cv_range_workbook": f"{lo_cv} to {hi_cv}" if hi_cv else f"{lo_cv}+",
                      "se_range_workbook": f"{lo_se} to {hi_se}" if hi_se else f"{lo_se}+",
                      "observed_cv_range": rng(code, "CV"), "observed_se_range": rng(code, "SE"),
                      "cells": int((numeric.quality == code).sum()), "has_numeric_estimate": True,
                      "usable": {"E": "with caution", "F": "no"}.get(code, "yes"),
                      "suppression_reason": "none (quality rating of sampling error)",
                      "sample_size_restriction": "not documented",
                      "recommended_treatment": {"E": "Use only with a caution flag; do not rank segments on E estimates alone",
                                                "F": "Keep the printed value but exclude from analysis (source says do not use)"}.get(code, "Use")})
    qrows += [
        {"marker": "..", "marker_type": "replaces the value (quality and error cells blank)",
         "official_meaning": "Not defined in the workbook or IMDB record 2941. Statistics Canada standard symbol: 'not available for a specific reference period'",
         "cells": int(len(dot)), "has_numeric_estimate": False, "usable": "no",
         "suppression_reason": (f"{dot_by_ctx.get('not_applicable_by_definition', 0)} cells not applicable to the segment by definition (INFERENCE); "
                                f"{dot_by_ctx.get('unavailable_reason_undocumented', 0)} cells unavailable for an undocumented reason (possibly small-sample suppression)"),
         "sample_size_restriction": "not documented",
         "recommended_treatment": "Treat as unavailable; never zero-fill or impute. Definitional cases are structurally not applicable"},
        {"marker": "…", "marker_type": "footnote symbol",
         "official_meaning": "'not applicable' (workbook footnote on every table sheet)",
         "cells": ellipsis_cells, "has_numeric_estimate": False, "usable": "not applicable",
         "suppression_reason": "none", "sample_size_restriction": "not applicable",
         "recommended_treatment": f"Defined but never used in a data cell ({ellipsis_footnotes} footnote lines)"},
        {"marker": "(blank)", "marker_type": "empty value cell",
         "official_meaning": "Not documented; all blanks are Table 50 region-row x region-column cells",
         "cells": int(len(blanks)), "has_numeric_estimate": False, "usable": "not applicable",
         "suppression_reason": "structurally omitted combination (a region cannot be cross-tabulated with itself)",
         "sample_size_restriction": "not applicable", "recommended_treatment": "Structural: not an estimate; never zero-fill"},
        {"marker": "0", "marker_type": "published numeric value",
         "official_meaning": "Displayed with one decimal; true zero and rounded zero cannot be distinguished",
         "cells": int(len(zeros)), "has_numeric_estimate": True, "usable": "yes (as published)",
         "suppression_reason": "none",
         "sample_size_restriction": f"{len(zeros_def)} zeros are definitional (e.g. Potential Borrowers' request rate)",
         "recommended_treatment": "Use as published (a value below 0.05 for percentages); exclude definitional rows from comparisons"},
        {"marker": "100 (definitional)", "marker_type": "published numeric value",
         "official_meaning": "Rejected borrowers requested financing by definition",
         "cells": int(len(hundreds_def)), "has_numeric_estimate": True, "usable": "not informative",
         "suppression_reason": "none", "sample_size_restriction": "not applicable",
         "recommended_treatment": "Exclude definitional rows from comparisons"},
    ]
    qmark = pd.DataFrame(qrows)
    write_csv(qmark, "quality_markers.csv")
    gate("every_quality_marker_classified", qmark.recommended_treatment.notna().all() and qmark.official_meaning.notna().all())
    print(qmark[["marker", "cells", "usable"]].to_string(index=False))

    # ---------------------------------------------------
    # 6. Estimate definitions + denominators
    # ---------------------------------------------------
    banner("6. ESTIMATE DEFINITIONS AND DENOMINATORS")
    all_row = cells[(cells.group == "All SMEs")]
    mstats = cells.assign(numeric_cells=cells.value_state.eq("numeric"), dotdot_cells=cells.value_state.eq("dotdot"),
                          blank_cells=cells.value_state.eq("blank"), quality_E=cells.quality.eq("E"),
                          quality_F=cells.quality.eq("F")).groupby("measure_id")[
        ["numeric_cells", "dotdot_cells", "blank_cells", "quality_E", "quality_F"]].sum()
    rows_out = []
    for r in meas.itertuples():
        tb = tables_by_no[r.table_no]
        ann = ANNOTATIONS.get((r.table_no, r.block_no))
        block_pct = all_row[(all_row.table_id == r.table_id) & (all_row.block_no == r.block_no)
                            & (all_row.estimate_type == "percentage_of_universe")].value
        block_sum = round(float(block_pct.sum()), 1) if len(block_pct) else None
        q_multiple = "multiple" in tb["t"]["questions"].lower()
        if ann:
            universe = ann["universe"]
            for key, uni in ann.get("measure_rules", []):
                if key in r.header_path.lower():
                    universe = uni
            status = "annotated"
            sums_100 = block_sum is not None and 99.5 <= block_sum <= 100.5
            multiple = ann["multiple"] or ("yes (block sums to %.1f in All SMEs row)" % block_sum if block_sum and block_sum > 100.5
                                           else "yes (question text)" if q_multiple
                                           else "no (block sums to %.1f in All SMEs row)" % block_sum if sums_100 else "not stated")
            exclusive = ann["exclusive"] or ("no" if block_sum and block_sum > 100.5
                                             else "likely yes (block sums to %.1f in All SMEs row)" % block_sum if sums_100 else "not stated")
            evidence = ann["evidence"]
        else:
            universe = "Unresolved at measure level; table notes: " + tb["t"]["notes"][:400]
            status = "unresolved_not_annotated"
            multiple = "yes (question text)" if q_multiple else "not stated"
            exclusive = "not stated"
            evidence = "Table not classified Core/Contextual; not annotated"
        mc = mstats.loc[r.measure_id]
        rows_out.append({**r._asdict(), "universe_or_denominator": universe, "denominator_status": status,
                         "denominator_evidence": evidence, "multiple_response": multiple,
                         "categories_exclusive": exclusive, "block_all_smes_percent_sum": block_sum,
                         "numeric_cells": int(mc.numeric_cells), "dotdot_cells": int(mc.dotdot_cells),
                         "blank_cells": int(mc.blank_cells), "quality_E": int(mc.quality_E), "quality_F": int(mc.quality_F)})
    est = pd.DataFrame(rows_out).drop(columns=["Index"])
    # Header anomalies (Table 38 row 14): mark measures whose value column sits under the stray label.
    t38 = tables_by_no["38"]["sh"]
    anomaly_found = {coord: clean(t38.value(int(coord[1:]), openpyxl.utils.cell.column_index_from_string(coord[0])) or "")
                     for (_, coord) in HEADER_ANOMALIES}
    gate("documented_header_anomalies_present_in_source",
         all(anomaly_found[coord] == txt for (_, coord), txt in HEADER_ANOMALIES.items()), str(anomaly_found))
    est["header_anomaly"] = [
        (f"stray header text '{HEADER_ANOMALIES[(r.table_no, r.value_col + '14')]}' in {r.value_col}14; "
         "row-13 label is the meaningful category") if (r.table_no, r.value_col + "14") in HEADER_ANOMALIES else None
        for r in est.itertuples()]
    write_csv(est, "estimate_definitions.csv")

    relevant = {k for k, v in RELEVANCE.items() if v[0] in ("Core", "Contextual")}
    rel_est = est[est.table_no.isin(relevant)]
    gate("every_relevant_measure_has_estimate_type_and_unit", rel_est[["estimate_type", "unit"]].notna().all().all())
    gate("every_relevant_measure_denominator_annotated", (rel_est.denominator_status == "annotated").all())
    gate("every_percentage_denominator_identified_or_marked_unresolved",
         est[est.unit_token == "%"].denominator_status.isin(["annotated", "unresolved_not_annotated"]).all())
    print(est.estimate_type.value_counts().to_string())

    # ---------------------------------------------------
    # 7. Dimension availability
    # ---------------------------------------------------
    banner("7. DIMENSION AVAILABILITY")
    cells["availability_dimension"] = [
        ("province" if gl == "province" else "sub_provincial_region" if gl == "sub_provincial_region"
         else "geography_region" if dc == "geography_region" else dc)
        for dc, gl in zip(cells.dimension_class, cells.geography_level)]
    dims = sorted(set(cells.availability_dimension))
    av_rows = []
    for tb in tables:
        tc = cells[cells.table_id == tb["table_id"]]
        for d in dims + ["revenue_band"]:
            dc = tc[tc.availability_dimension == d]
            n_meas = tc.measure_id.nunique()
            per_cat = dc.groupby("label").value_state.agg(lambda s: (s == "numeric").sum())
            num = dc[dc.value_state == "numeric"]
            av_rows.append({"table_id": tb["table_id"], "dimension": d,
                            "categories_in_stub": int(dc.label.nunique()),
                            "categories_with_any_estimate": int((per_cat > 0).sum()),
                            "categories_with_estimate_in_every_measure": int((per_cat == n_meas).sum()),
                            "numeric_cells": int(len(num)),
                            "cells_quality_A_to_D": int(num.quality.isin(list("ABCD")).sum()),
                            "cells_quality_E": int((num.quality == "E").sum()),
                            "cells_quality_F": int((num.quality == "F").sum()),
                            "cells_dotdot": int((dc.value_state == "dotdot").sum()),
                            "cells_blank": int((dc.value_state == "blank").sum())})
    avail = pd.DataFrame(av_rows)
    write_csv(avail, "dimension_availability.csv")
    # Tidy cells (with dimension classes and '..' contexts) for later phases; disposable, gitignored.
    cells.drop(columns=["value_number_format"]).to_parquet(TIDY_CACHE, index=False)
    # Reliability of Core-table estimates by key dimension (share of printed estimates by quality).
    core_ids = [f"ISED_T{int(k):02d}" for k, v in RELEVANCE.items() if v[0] == "Core"]
    key_dims = ["employment_size", "geography_region", "province", "geography_cma", "industry_naics_group"]
    cr = avail[avail.table_id.isin(core_ids) & avail.dimension.isin(key_dims) & (avail.numeric_cells > 0)].copy()
    for col, src in (("pct_A_to_D", "cells_quality_A_to_D"), ("pct_E", "cells_quality_E"), ("pct_F", "cells_quality_F")):
        cr[col] = (100 * cr[src] / cr.numeric_cells).round(1)
    core_reliability = {f"{r.table_id}|{r.dimension}": {"A_to_D": r.pct_A_to_D, "E": r.pct_E, "F": r.pct_F, "estimates": int(r.numeric_cells)}
                        for r in cr.itertuples()}
    rb = avail[avail.dimension == "revenue_band"]
    gate("revenue_band_dimension_absent_in_every_table", (rb.categories_in_stub == 0).all())

    # ---------------------------------------------------
    # 8. Reconciliation
    # ---------------------------------------------------
    banner("8. RECONCILIATION")
    # pandas stores unavailable values as NaN; map them back to None so they are skipped, not compared.
    V = {(r.table_no, r.measure_idx, r.group, r.label): (None if pd.isna(r.value) else r.value) for r in cells.itertuples()}

    def mids(tno, block=None, contains=None, not_contains=None):
        m = meas[meas.table_no == tno]
        if block is not None:
            m = m[m.block_no == block]
        if contains:
            m = m[m.header_path.str.lower().str.contains(contains.lower(), regex=False)]
        if not_contains:
            m = m[~m.header_path.str.lower().str.contains(not_contains.lower(), regex=False)]
        return list(m.measure_idx)

    def stub_rows(tno, exclude_definitional=True):
        return [(g, l) for g, l in tables_by_no[tno]["stub_keys"]
                if not (exclude_definitional and cls_map[(g, l)][0] in DEFINITIONAL_CLASSES)]

    def rows_of(tno, cls=None, labels=None):
        return [(g, l) for g, l in tables_by_no[tno]["stub_keys"]
                if (cls is None or cls_map[(g, l)][0] == cls) and (labels is None or l in labels)]

    ALLK = ("All SMEs", "1 to 499 employees")
    TOP_REGIONS = ["Atlantic", "Quebec", "Ontario", "Prairies", "British Columbia and Territories"]
    W, N, P = 2, 1, 3            # Table 49 measure indices (block order): respondents, weighted count, percentage
    size_rows = rows_of("49", "employment_size")
    region_rows = [("Region", l) for l in TOP_REGIONS]
    ind_rows = rows_of("49", "industry_naics_group")

    def parent_children(tno, mi, parent, children):
        return [("all rows", V.get((tno, mi, *parent)), [V.get((tno, mi, *c)) for c in children])]

    def rowwise(tno, parts, total=None, parent_mi=None, exclude_definitional=True):
        out = []
        for g, l in stub_rows(tno, exclude_definitional):
            kids = [V.get((tno, mi, g, l)) for mi in parts]
            par = V.get((tno, parent_mi, g, l)) if parent_mi else total
            out.append((f"{g} | {l}", par, kids))
        return out

    t50_size = {meas[(meas.table_no == "50") & (meas.measure_idx == i)].header_path.iloc[0].split(" > ")[-1]: i
                for i in mids("50", block=1)}
    t50_reg = {meas[(meas.table_no == "50") & (meas.measure_idx == i)].header_path.iloc[0].split(" > ")[-1]: i
               for i in mids("50", block=2)}
    T50_TOP = ["Atlantic", "Quebec", "Ontario", "Prairies", "British Columbia"]
    T50_TO_T49_REGION = {"Atlantic": "Atlantic", "Quebec": "Quebec", "Ontario": "Ontario", "N. Ontario": "Northern Ontario",
                         "S. Ontario": "Southern Ontario", "Prairies": "Prairies", "Manitoba": "Manitoba",
                         "Saskatchewan": "Saskatchewan", "Alberta": "Alberta", "British Columbia": "British Columbia and Territories"}
    t50_rows_region = rows_of("50", "geography_region")
    t50_rows_ind = rows_of("50", "industry_naics_group")
    t50_top_region_rows = [r for r in t50_rows_region if r[1] in TOP_REGIONS]

    def t50_vs_t49_rows(rows):
        return [(f"{g} | {l}", V.get(("49", W, "Industries" if g == "Industry" else g, l)),
                 [V.get(("50", i, g, l)) for i in t50_size.values()]) for g, l in rows]

    def t50_regions_vs_t49(rows):
        return [(f"{g} | {l}", V.get(("49", W, "Industries", l)), [V.get(("50", t50_reg[c], g, l)) for c in T50_TOP])
                for g, l in rows]

    def t50_cols(rows, band_label):
        return (band_label, V.get(("49", W, "Employment Size", band_label)),
                [V.get(("50", t50_size[band_label], g, l)) for g, l in rows])

    obstacles = sorted({p.split(" > ")[1] for p in meas[meas.table_no == "20"].header_path})

    def t20():
        out = []
        for ob in obstacles:
            parts = [i for i in mids("20") if meas[(meas.table_no == "20") & (meas.measure_idx == i)].header_path.iloc[0].split(" > ")[1] == ob]
            out += [(f"{ob} | {c}", 100.0, k) for c, _, k in rowwise("20", parts, total=100.0)]
        return out

    DECL = dict(exclusive=True, same_population=True, same_unit=True, same_period=True,
                quality_treatment="All published values used; E and F components counted in the result")
    specs = [
        dict(test_id="R_T49_SIZE_W", tables="T49", unit="count", desc="Weighted count: four size bands sum to All SMEs", exhaustive="declared (1-499 partition)",
             comps=lambda: parent_children("49", W, ALLK, size_rows)),
        dict(test_id="R_T49_SIZE_N", tables="T49", unit="count", desc="Respondents: four size bands sum to All SMEs", exhaustive="declared (1-499 partition)",
             comps=lambda: parent_children("49", N, ALLK, size_rows)),
        dict(test_id="R_T49_REGION_W", tables="T49", unit="count", desc="Weighted count: five regions sum to All SMEs", exhaustive="declared (regions cover Canada)",
             comps=lambda: parent_children("49", W, ALLK, region_rows)),
        dict(test_id="R_T49_REGION_N", tables="T49", unit="count", desc="Respondents: five regions sum to All SMEs", exhaustive="declared",
             comps=lambda: parent_children("49", N, ALLK, region_rows)),
        dict(test_id="R_T49_ONTARIO_W", tables="T49", unit="count", desc="Weighted count: Northern + Southern Ontario = Ontario", exhaustive="declared",
             comps=lambda: parent_children("49", W, ("Region", "Ontario"), [("Region", "Northern Ontario"), ("Region", "Southern Ontario")])),
        dict(test_id="R_T49_PRAIRIES_W", tables="T49", unit="count", desc="Weighted count: Manitoba + Saskatchewan + Alberta = Prairies", exhaustive="declared",
             comps=lambda: parent_children("49", W, ("Region", "Prairies"), [("Region", "Manitoba"), ("Region", "Saskatchewan"), ("Region", "Alberta")])),
        dict(test_id="R_T49_RURAL_URBAN_W", tables="T49", unit="count", desc="Weighted count: rural + urban = All SMEs", exhaustive="tested",
             comps=lambda: parent_children("49", W, ALLK, rows_of("49", "rural_urban"))),
        dict(test_id="R_T49_LEGAL_W", tables="T49", unit="count", desc="Weighted count: incorporated + unincorporated = All SMEs", exhaustive="tested",
             comps=lambda: parent_children("49", W, ALLK, [("Legal Status", "Incorporated"), ("Legal Status", "Unincorporated")])),
        dict(test_id="R_T49_LEGAL_N", tables="T49", unit="count", desc="Respondents: incorporated + unincorporated = All SMEs", exhaustive="tested",
             comps=lambda: parent_children("49", N, ALLK, [("Legal Status", "Incorporated"), ("Legal Status", "Unincorporated")])),
        dict(test_id="R_T49_AGE_W", tables="T49", unit="count", desc="Weighted count: four business-age classes = All SMEs", exhaustive="tested",
             comps=lambda: parent_children("49", W, ALLK, rows_of("49", "business_age"))),
        dict(test_id="R_T49_INDUSTRY_W", tables="T49", unit="count", desc="Weighted count: ten NAICS-based industry groups = All SMEs (special aggregates excluded: they overlap sectors)", exhaustive="tested",
             comps=lambda: parent_children("49", W, ALLK, ind_rows)),
        dict(test_id="R_T49_INDUSTRY_N", tables="T49", unit="count", desc="Respondents: ten industry groups = All SMEs", exhaustive="tested",
             comps=lambda: parent_children("49", N, ALLK, ind_rows)),
        dict(test_id="R_T49_SHARE", tables="T49", unit="share", desc="Percentage = 100 x weighted count / All-SMEs weighted count, every row", exhaustive="not applicable",
             comps=lambda: [(f"{g} | {l}", V.get(("49", P, g, l)),
                             [100.0 * V[("49", W, g, l)] / V[("49", W, *ALLK)]] if V.get(("49", W, g, l)) is not None else [None])
                            for g, l in stub_rows("49", exclude_definitional=False)]),
        dict(test_id="R_T50_ROWS_VS_T49", tables="T50,T49", unit="count", desc="Each Table 50 region/industry row: four size columns = Table 49 weighted count of the same segment", exhaustive="declared (1-499 partition)",
             comps=lambda: t50_vs_t49_rows(t50_rows_region + t50_rows_ind)),
        dict(test_id="R_T50_INDUSTRY_BY_REGION_VS_T49", tables="T50,T49", unit="count", desc="Each Table 50 industry row: Atlantic + Quebec + Ontario + Prairies + 'British Columbia' columns = Table 49 industry count (tests whether 'British Columbia' includes the territories)", exhaustive="tested",
             comps=lambda: t50_regions_vs_t49(t50_rows_ind)),
        dict(test_id="R_T50_REGION_COLS_VS_T49", tables="T50,T49", unit="count", desc="Each Table 50 region column: ten industry rows = Table 49 count of that region ('British Columbia' column vs 'British Columbia and Territories' row)", exhaustive="declared (industry groups tested exhaustive)",
             comps=lambda: [(c, V.get(("49", W, "Region", T50_TO_T49_REGION[c])), [V.get(("50", t50_reg[c], g, l)) for g, l in t50_rows_ind]) for c in t50_reg]),
        dict(test_id="R_T50_ONTARIO_COLS", tables="T50", unit="count", desc="Industry rows: N. Ontario + S. Ontario columns = Ontario column", exhaustive="declared",
             comps=lambda: [(f"{g} | {l}", V.get(("50", t50_reg["Ontario"], g, l)), [V.get(("50", t50_reg[c], g, l)) for c in ("N. Ontario", "S. Ontario")]) for g, l in t50_rows_ind]),
        dict(test_id="R_T50_PRAIRIE_COLS", tables="T50", unit="count", desc="Industry rows: Manitoba + Saskatchewan + Alberta columns = Prairies column", exhaustive="declared",
             comps=lambda: [(f"{g} | {l}", V.get(("50", t50_reg["Prairies"], g, l)), [V.get(("50", t50_reg[c], g, l)) for c in ("Manitoba", "Saskatchewan", "Alberta")]) for g, l in t50_rows_ind]),
        dict(test_id="R_T50_SIZE_BY_REGION_VS_T49", tables="T50,T49", unit="count", desc="Each size column: five region rows = Table 49 size-band count", exhaustive="declared",
             comps=lambda: [t50_cols(t50_top_region_rows, b) for b in t50_size]),
        dict(test_id="R_T50_SIZE_BY_INDUSTRY_VS_T49", tables="T50,T49", unit="count", desc="Each size column: ten industry rows = Table 49 size-band count", exhaustive="tested",
             comps=lambda: [t50_cols(t50_rows_ind, b) for b in t50_size]),
        dict(test_id="R_T01_REQUEST_COMPLEMENT", parent_constant=True, tables="T01", unit="percent", desc="Requested + did not request external financing = 100, every non-definitional row", exhaustive="declared",
             comps=lambda: rowwise("1", mids("1", 1) + mids("1", 2), total=100.0)),
        dict(test_id="R_T01_REASONS", parent_constant=True, tables="T01", unit="percent", desc="Six reasons for not requesting = 100, every non-definitional row", exhaustive="tested",
             comps=lambda: rowwise("1", mids("1", 3), total=100.0)),
        dict(test_id="R_T03_FULL_PARTIAL", tables="T03", unit="percent", desc="Full + partial approval = overall approval rate", exhaustive="declared",
             comps=lambda: rowwise("3", mids("3", 2, "full amount") + mids("3", 2, "partial amount"), parent_mi=mids("3", 2, "overall approval")[0])),
        dict(test_id="R_T03_OUTCOMES", parent_constant=True, tables="T03", unit="percent", desc="Overall approval + rejected + under review + withdrawn = 100", exhaustive="tested",
             comps=lambda: rowwise("3", mids("3", 2, "overall approval") + mids("3", 2, "rejected") + mids("3", 2, "under review") + mids("3", 2, "withdrawn"), total=100.0)),
        dict(test_id="R_T09_FULL_PARTIAL", tables="T09", unit="percent", desc="Lease: full + partial approval = overall approval rate", exhaustive="declared",
             comps=lambda: rowwise("9", mids("9", 2, "full amount") + mids("9", 2, "partial amount"), parent_mi=mids("9", 2, "overall approval")[0])),
        dict(test_id="R_T16_ESTABLISHED", parent_constant=True, tables="T16", unit="percent", desc="Four establishment-year classes = 100", exhaustive="tested",
             comps=lambda: rowwise("16", mids("16", 1, not_contains="average"), total=100.0)),
        dict(test_id="R_T16_SELLING", parent_constant=True, tables="T16", unit="percent", desc="Four years-selling classes = 100", exhaustive="tested",
             comps=lambda: rowwise("16", mids("16", 2, not_contains="average"), total=100.0)),
        dict(test_id="R_T19_PAST", parent_constant=True, tables="T19", unit="percent", desc="Seven past-growth classes = 100", exhaustive="tested",
             comps=lambda: rowwise("19", mids("19", 1, "over the last three years"), total=100.0)),
        dict(test_id="R_T19_EXPECTED", parent_constant=True, tables="T19", unit="percent", desc="Seven expected-growth classes = 100", exhaustive="tested",
             comps=lambda: rowwise("19", mids("19", 1, "over the next three years"), total=100.0)),
        dict(test_id="R_T20_SEVERITY", parent_constant=True, tables="T20", unit="percent", desc="Each obstacle: four severity levels = 100", exhaustive="tested",
             comps=t20),
        dict(test_id="R_T38_INTENT", parent_constant=True, tables="T38", unit="percent", desc="Do not intend + intend + do not know = 100", exhaustive="tested",
             comps=lambda: rowwise("38", mids("38", 1) + mids("38", 2) + mids("38", 4), total=100.0)),
        dict(test_id="R_T38_INTENTION_TYPE", parent_constant=True, tables="T38", unit="percent", desc="Six intention types = 100 among SMEs intending to exit", exhaustive="tested",
             comps=lambda: rowwise("38", mids("38", 3), total=100.0)),
    ]
    rec_rows = []
    for sp in specs:
        comps = sp["comps"]()
        n_ok = n_out = n_skip = 0
        max_diff, worst = 0.0, None
        for ctx_label, parent, kids in comps:
            if parent is None or any(k is None for k in kids):
                n_skip += 1
                continue
            terms = len(kids) + (0 if sp.get("parent_constant") else 1)
            tol = {"count": 0.5 * terms, "percent": 0.05 * terms, "share": 0.0501}[sp["unit"]]
            diff = sum(kids) - parent
            if abs(diff) <= tol + 1e-9:
                n_ok += 1
            else:
                n_out += 1
            if abs(diff) > abs(max_diff):
                max_diff, worst = diff, ctx_label
        status = ("not_testable" if n_ok + n_out == 0 else "consistent" if n_out == 0 else "finding")
        rec_rows.append({"test_id": sp["test_id"], "tables": sp["tables"], "description": sp["desc"], "unit": sp["unit"],
                         "comparisons": n_ok + n_out, "within_tolerance": n_ok, "outside_tolerance": n_out,
                         "skipped_unavailable": n_skip, "max_abs_difference": round(abs(max_diff), 4),
                         "max_signed_difference": round(max_diff, 4), "worst_case": worst,
                         "tolerance_rule": {"count": "0.5 x rounded terms", "percent": "0.05 x rounded terms", "share": "0.05 (one-decimal rounding)"}[sp["unit"]],
                         "status": status, "categories_exhaustive": sp["exhaustive"], **DECL})
    # Cross-source checks (text constants vs workbook).
    t49_incorp = V.get(("49", P, "Legal Status", "Incorporated"))
    ta_cov = re.search(r"represents ([\d.]+)% of all SMEs", " ".join(tables_by_no["A"]["t"]["footnotes"]))
    ta_cov = float(ta_cov.group(1)) if ta_cov else None
    xsrc = [
        ("R_X_RESPONDENTS_NOTES", "Table 49 All-SMEs respondents = responses stated in the Notes sheet", V.get(("49", N, *ALLK)), notes_responses),
        ("R_X_SAMPLE_IMDB", "Notes sample size = IMDB record 2941 final unduplicated sample", notes_sample, IMDB_SAMPLE_SIZE),
    ]
    for tid_, desc, a, b in xsrc:
        rec_rows.append({"test_id": tid_, "tables": "cross-source", "description": desc, "unit": "count", "comparisons": 1,
                         "within_tolerance": int(a == b), "outside_tolerance": int(a != b), "skipped_unavailable": 0,
                         "max_abs_difference": abs((a or 0) - (b or 0)), "max_signed_difference": (a or 0) - (b or 0), "worst_case": f"{a} vs {b}",
                         "tolerance_rule": "exact", "status": "consistent" if a == b else "finding", "categories_exhaustive": "not applicable", **DECL})
    rec_rows.append({"test_id": "R_X_TABLEA_COVERAGE", "tables": "TA,T49", "description": "Table A coverage (incorporated with complete tax data) <= Table 49 incorporated share",
                     "unit": "percent", "comparisons": 1, "within_tolerance": int(ta_cov <= t49_incorp), "outside_tolerance": int(ta_cov > t49_incorp),
                     "skipped_unavailable": 0, "max_abs_difference": round(t49_incorp - ta_cov, 1), "max_signed_difference": round(ta_cov - t49_incorp, 1),
                     "worst_case": f"{ta_cov} vs {t49_incorp}", "tolerance_rule": "inequality", "status": "consistent" if ta_cov <= t49_incorp else "finding",
                     "categories_exhaustive": "not applicable", **DECL})
    recon = pd.DataFrame(rec_rows)
    write_csv(recon, "reconciliation_results.csv")
    gate("reconciliations_use_only_declared_compatible_categories",
         recon[["exclusive", "same_population", "same_unit", "same_period"]].all().all())
    gate("every_reconciliation_executed", (recon.status != "not_testable").all())
    print(recon[["test_id", "comparisons", "within_tolerance", "outside_tolerance", "skipped_unavailable", "max_abs_difference", "status"]].to_string(index=False))

    # ---------------------------------------------------
    # 9. Table inventory + project relevance
    # ---------------------------------------------------
    banner("9. TABLE INVENTORY AND RELEVANCE")
    inv, relv = [], []
    for tb in tables:
        t, tno, tid = tb["t"], tb["table_no"], tb["table_id"]
        tc = cells[cells.table_id == tid]
        te = est[est.table_id == tid]
        dims_present = sorted(set(tc[tc.value_state == "numeric"].availability_dimension))
        stub_first, stub_last = t["stub"][0]["row"], t["stub"][-1]["row"]
        last_foot = max(r for (r, c) in tb["sh"].nonempty)
        yrs, rels = years_and_periods(" ".join([t["title"], t["questions"], t["notes"]]))
        cls = RELEVANCE.get(tno, ("Out of scope",))[0]
        inv.append({
            "table_id": tid, "sheet": tb["sheet"], "sheet_order": tb["sheet_order"],
            "table_range": f"A1:{col_letter(t['last_col'])}{last_foot}",
            "data_range": f"D{stub_first}:{col_letter(t['last_col'])}{stub_last}",
            "title": t["title"], "index_title": index_titles.get(tno),
            "index_title_matches_sheet": clean(index_titles.get(tno, "")).lower() == clean(t["title"]).lower(),
            "survey_questions": ";".join(sorted(set(re.findall(r"Q\d+", t["questions"])), key=lambda q: int(q[1:]))),
            "universe_notes": t["notes"], "multiple_response_in_questions": "multiple" in t["questions"].lower(),
            "estimate_types": ";".join(sorted(set(te.estimate_type))), "units": ";".join(sorted(set(te.unit))),
            "blocks": te.block_id.nunique(), "measures": len(te), "stub_rows": len(t["stub"]),
            "reference_period": "2023 (survey reference year)" + (f"; other periods in questions: {rels}" if rels else ""),
            "years_mentioned": yrs,
            "geography": ";".join(d for d in dims_present if d in ("all_smes", "geography_region", "province", "sub_provincial_region", "geography_cma")),
            "industry": ";".join(d for d in dims_present if d.startswith("industry")) + (" (NAICS 2022 per IMDB)" if any(d.startswith("industry") for d in dims_present) else ""),
            "size_bands": "employment_size (1-4, 5-19, 20-99, 100-499)" if "employment_size" in dims_present or tno == "50" else "",
            "other_segmentation": ";".join(d for d in dims_present if d not in ("all_smes", "geography_region", "province", "sub_provincial_region", "geography_cma", "employment_size") and not d.startswith("industry")),
            "sample_size_field": "yes (Number of Respondents)" if tno == "49" else "no (see Table 49)",
            "total_subtotal_structure": ("Rows: All SMEs total; one-way groups; Ontario north/south and Prairie provinces nest in regions; special industry aggregates overlap sectors"
                                         if tno != "50" else "Rows: regions and industry groups; columns: size bands and regions; region-by-region cells blank"),
            "numeric_cells": int((tc.value_state == "numeric").sum()), "dotdot_cells": int((tc.value_state == "dotdot").sum()),
            "blank_cells": int((tc.value_state == "blank").sum()), "quality_E": int((tc.quality == "E").sum()),
            "quality_F": int((tc.quality == "F").sum()), "footnotes": " | ".join(t["footnotes"]),
            "relevance": cls,
        })
        if cls in ("Core", "Contextual"):
            _, concept, what, bq, decision = RELEVANCE[tno]
            relv.append({"table_id": tid, "title": t["title"], "relevance": cls, "rationale": what,
                         "business_question": bq, "decision_supported": decision, "concept": concept,
                         "usable_grain": ("Two-way weighted counts: size x region (5 regions + provinces), size x industry (10 groups), industry x region" if tno == "50"
                                          else "National and size-band level only (CMA and industry estimates are largely E/F)" if tno == "5"
                                          else "One-way segments (as Table 1-49 stub)" if tno != "A" else "One-way segments; incorporated SMEs only"),
                         "comparable_with_0722_1097": ("Enterprise unit as in 0722 (not 1097's locations); narrower SME population; "
                                                       "size bands identical to 0722; CMAs by name only; industry groups coarser than both" if tno in ("49", "50")
                                                       else "Not directly: survey prevalence of financing behaviour has no counterpart in 0722 or 1097; can be applied only as a segment attribute"),
                         "cautions": CAUTION_STD + (" Table A covers incorporated SMEs with complete tax data (89.8%)." if tno == "A" else "")
                                     + (f" Deal-size estimates are weak below the national level: at CMA level {core_reliability['ISED_T05|geography_cma']['A_to_D']}% are A-D "
                                        f"and {core_reliability['ISED_T05|geography_cma']['F']}% are F; by industry {core_reliability['ISED_T05|industry_naics_group']['F']}% are F."
                                        if tno == "5" else "")
                                     + (f" {core_reliability['ISED_T04|geography_cma']['F']}% of CMA-level estimates are F." if tno == "4" else ""),
                         "reference_year_alignment": ALIGN_STD})
        else:
            relv.append({"table_id": tid, "title": t["title"], "relevance": "Out of scope",
                         "rationale": OUT_OF_SCOPE_REASON[tno], "business_question": None, "decision_supported": None,
                         "concept": None, "usable_grain": None, "comparable_with_0722_1097": None, "cautions": None,
                         "reference_year_alignment": None})
    tinv = pd.DataFrame(inv)
    write_csv(tinv, "table_inventory.csv")
    relevance = pd.DataFrame(relv)
    write_csv(relevance, "project_relevance.csv")
    write_csv(wbinv, "workbook_inventory.csv")
    rel_counts = relevance.relevance.value_counts().to_dict()
    gate("every_table_classified_for_relevance", relevance.rationale.notna().all() and len(relevance) == len(tables))
    print(rel_counts)

    # ---------------------------------------------------
    # 10. Compatibility observations (diagnostic)
    # ---------------------------------------------------
    banner("10. COMPATIBILITY WITH 33-10-0722 AND 33-10-1097 (DIAGNOSTIC)")
    ised_cmas = [l for g, l in tables_by_no["1"]["stub_keys"] if g == "CMA Level"]
    g0722 = pd.read_csv(GEO_0722)
    g1097 = pd.read_csv(GEO_1097, dtype={"geo_code": str})
    c0722 = g0722[g0722.geo_level == "CMA"].geo.tolist()
    c1097 = g1097[g1097.geo_level == "CMA"].geo.tolist()
    cma_rows = []
    for c in ised_cmas:
        cma_rows.append({"ised_cma_label": c,
                         "exact_match_1097": c in c1097, "normalized_match_1097": next((x for x in c1097 if norm_place(x) == norm_place(c)), None),
                         "exact_match_0722": c in c0722, "normalized_match_0722": next((x for x in c0722 if norm_place(x) == norm_place(c)), None)})
    cma_diag = pd.DataFrame(cma_rows)
    write_csv(cma_diag, "cma_label_diagnostic.csv")
    unmatched = cma_diag[cma_diag.normalized_match_1097.isna()].ised_cma_label.tolist()
    ised_bands = [l for g, l in tables_by_no["1"]["stub_keys"] if g == "Employment Size"]
    bands_0722 = pd.read_csv(ROOT / "data" / "validation" / "33100722" / "availability_by_employment_size.csv").employment_size.tolist()
    bands_equal_0722 = set(ised_bands) <= set(bands_0722)
    comp = pd.DataFrame([
        ("statistical unit", "Enterprise (IMDB record 2941)", "Enterprise (0722 Note 2)", "Statistical location (1097 Note 1)", "Same unit as 0722; not 1097", "FACT: SOURCE"),
        ("SME definition", "1-499 employees and >= $30,000 revenue; excludes NAICS 22, 52, 55, 61, 91 and 5321, 5324, 6214-6223, 6242, 814110; excludes non-profits, joint ventures, government, subsidiaries, financing and leasing companies",
         "All employer enterprises (incl. 500+); business-sector total T004", "All employer locations, all industries incl. 91", "ISED population is a filtered subset; totals are not comparable without scope alignment", "FACT: SOURCE"),
        ("employment-size bands", "; ".join(ised_bands), "1-4; 5-19; 20-99; 100-499; 500+", "1-4; 5-9; 10-19; 20-49; 50-99; 100-199; 200-499; 500+",
         "ISED band labels are identical to 0722's" if bands_equal_0722 else "Band labels differ from 0722", "FACT: DATA"),
        ("industry classification", "NAICS 2022 (IMDB); 10 groups, some combined (11+21; 51+53+56+62+71); overlapping special aggregates", "Version not stated; 18 sector members", "NAICS 2022, to 4-digit",
         "Needs a group-level crosswalk (Phase E); ISED groups are coarser", "FACT: SOURCE / DATA"),
        ("geography", f"Canada; 5 regions; QC, ON, MB, SK, AB; N/S Ontario; {len(ised_cmas)} CMAs by name only (no codes; vintage not stated)",
         "Canada, provinces, territories, 35 CMAs (2016 DGUIDs)", "41 CMAs + CSDs (2021 DGUIDs); no provinces",
         f"CMA labels: {int(cma_diag.normalized_match_1097.notna().sum())} of {len(ised_cmas)} match 1097 after normalization; no codes, so any join needs a name crosswalk (Phase E)", "FACT: DATA"),
        ("reference period", "Calendar year 2023 (collection Feb-Jun 2024; Table A: 2023 CRA data)", "Monthly 2015-01 to 2026-05", "December 2025 snapshot",
         "Annual 2023 vs monthly vs Dec 2025: align only for context", "FACT: SOURCE"),
        ("estimate type", "Weighted survey estimates (percentages, dollars, counts) with CV/SE quality codes", "Experimental administrative estimates, seasonally adjusted", "Register counts",
         "Survey sampling error has no counterpart in 0722/1097", "FACT: SOURCE"),
        ("population coverage", "863,533 weighted SMEs (All SMEs); a filtered population", "Active enterprises (all sizes)", "Employer locations",
         "Magnitude comparison deferred to Phase D", "FACT: DATA"),
    ], columns=["dimension", "ised_2023", "table_33100722", "table_33101097", "observation", "evidence_label"])
    write_csv(comp, "compatibility_observations.csv")
    print(f"  CMA labels matching 1097 after normalization: {int(cma_diag.normalized_match_1097.notna().sum())} of {len(ised_cmas)}; unmatched: {unmatched}")

    # ---------------------------------------------------
    # 11. Summary facts + documentation claims
    # ---------------------------------------------------
    w_all = V[("49", W, *ALLK)]
    n_all = V[("49", N, *ALLK)]
    facts = {
        "worksheets": n_sheets, "table_sheets": len(tables), "front_sheets": n_sheets - len(tables),
        "blocks": int(meas.block_id.nunique()), "measures": int(len(meas)),
        "value_cells": int(len(cells)), "numeric_values": int((cells.value_state == "numeric").sum()),
        "dotdot_values": int(len(dot)), "blank_values": int(len(blanks)),
        "dotdot_definitional": int(dot_by_ctx.get("not_applicable_by_definition", 0)),
        "dotdot_undocumented": int(dot_by_ctx.get("unavailable_reason_undocumented", 0)),
        "quality_coded_values": n_coded, "quality_E": e_printed, "quality_F": f_printed,
        "quality_A_to_D": int(numeric.quality.isin(list("ABCD")).sum()),
        "workbook_threshold_agreement": agree["workbook_notes"], "imdb_threshold_agreement": agree["imdb_record"],
        "weighted_smes": int(w_all), "respondents": int(n_all),
        "notes_sample": notes_sample, "notes_responses": notes_responses, "notes_response_rate": notes_rate,
        "stub_rows_main": len(ref_stub), "stub_rows_t50": len(tables_by_no["50"]["stub_keys"]),
        "cmas": len(ised_cmas), "industry_groups": int(dim_counts.get("industry_naics_group", 0)),
        "industry_special_aggregates": int(dim_counts.get("industry_special_aggregate", 0)),
        "size_bands": int(dim_counts.get("employment_size", 0)),
        "provinces_individual": int(geo_counts.get("province", 0)),
        "owner_demographic_rows": int(main_cat.owner_demographic.sum()),
        "defined_names": len(pkg["defined_names"]), "merged_ranges": int(wbinv.merged_ranges.sum()),
        "core_tables": int(rel_counts.get("Core", 0)), "contextual_tables": int(rel_counts.get("Contextual", 0)),
        "out_of_scope_tables": int(rel_counts.get("Out of scope", 0)),
        "reconciliation_tests": int(len(recon)), "reconciliation_consistent": int((recon.status == "consistent").sum()),
        "reconciliation_findings": int((recon.status == "finding").sum()),
        "cma_labels_matching_1097": int(cma_diag.normalized_match_1097.notna().sum()),
        "table_a_coverage_pct": ta_cov, "t49_incorporated_pct": t49_incorp,
        "percentage_measures": int((meas.unit_token == "%").sum()),
        "conditional_format_rules": len(cf_rules), "conditional_format_rules_x": len(cf_x_rules), "x_cells": x_cells,
        "custom_views": pkg["custom_workbook_views"], "index_title_mismatches": int((~tinv.index_title_matches_sheet).sum()),
        "zeros_published": int(len(zeros)), "zeros_definitional": int(len(zeros_def)),
        "annotated_measures": int((est.denominator_status == "annotated").sum()),
        "nonprofit_se_weighted": int(V[("49", W, "Social Enterprises", "Not-for-profit")] + V[("49", W, "Social Enterprises", "Charity")]),
        "nonprofit_se_pct_of_all": round(100 * (V[("49", W, "Social Enterprises", "Not-for-profit")] + V[("49", W, "Social Enterprises", "Charity")]) / w_all, 2),
        "t49_share_rows_matching": int(recon.loc[recon.test_id == "R_T49_SHARE", "within_tolerance"].iloc[0]),
        "t49_share_rows_other_base": int(recon.loc[recon.test_id == "R_T49_SHARE", "outside_tolerance"].iloc[0]),
        "t05_cma_pct_A_to_D": core_reliability["ISED_T05|geography_cma"]["A_to_D"],
        "t05_cma_pct_F": core_reliability["ISED_T05|geography_cma"]["F"],
        "t05_industry_pct_F": core_reliability["ISED_T05|industry_naics_group"]["F"],
        "t01_cma_pct_A_to_D": core_reliability["ISED_T01|geography_cma"]["A_to_D"],
        "t03_cma_pct_A_to_D": core_reliability["ISED_T03|geography_cma"]["A_to_D"],
        "t04_cma_pct_F": core_reliability["ISED_T04|geography_cma"]["F"],
    }

    def fmt(n):
        return f"{n:,}" if isinstance(n, int) else str(n)

    F = {k: fmt(v) for k, v in facts.items()}
    # Each claim is anchored to its label as written in the Phase C section, so it cannot match by coincidence.
    doc_claims = {
        "worksheets": f"**{F['worksheets']}** worksheets", "table_sheets": f"**{F['table_sheets']}** table sheets",
        "blocks": f"**{F['blocks']}** in total", "measures": f"Blocks contain **{F['measures']}** measures",
        "numeric_values": f"There are {F['numeric_values']} numeric values",
        "quality_coded_values": f"**{F['quality_coded_values']}** of them carry a quality code",
        "quality_A_to_D": f"| A–D | **{F['quality_A_to_D']}** |", "quality_E": f"| E | **{F['quality_E']}** |",
        "quality_F": f"| F | **{F['quality_F']}** |", "dotdot_values": f"| `..` | **{F['dotdot_values']}** |",
        "blank_values": f"| blank | **{F['blank_values']}** |", "zeros_published": f"| 0 | **{F['zeros_published']}** |",
        "dotdot_definitional": f"**{F['dotdot_definitional']}** cells fall in segments",
        "dotdot_undocumented": f"**{F['dotdot_undocumented']}** cells are unavailable",
        "imdb_threshold_agreement": f"only **{F['imdb_threshold_agreement']}**",
        "weighted_smes": f"**{F['weighted_smes']}** SMEs in the target population",
        "respondents": f"(**{F['respondents']}**) are sample support",
        "notes_sample": f"**{F['notes_sample']}** enterprises", "notes_response_rate": f"gives **{facts['notes_response_rate']}%**",
        "stub_rows_main": f"(**{F['stub_rows_main']}** rows)", "cmas": f"| CMA | **{F['cmas']}** |",
        "industry_groups": f"| Industry, NAICS-based | **{F['industry_groups']}** |",
        "relevance": f"**{F['core_tables']}** Core, **{F['contextual_tables']}** Contextual and **{F['out_of_scope_tables']}** out of scope",
        "reconciliation_tests": f"There are **{F['reconciliation_tests']}** tests",
        "reconciliation_results": f"**{F['reconciliation_consistent']}** consistent and **{F['reconciliation_findings']}** finding",
        "cma_labels_matching_1097": f"**{F['cma_labels_matching_1097']}** of {F['cmas']} CMA labels",
        "defined_names": f"all **{F['defined_names']}** are hidden", "custom_views": f"**{F['custom_views']}** saved custom views",
        "merged_ranges": f"**Merged ranges:** **{F['merged_ranges']}**",
        "conditional_format_rules": f"**{F['conditional_format_rules']}** rules",
        "percentage_measures": f"**{F['percentage_measures']}** are percentages",
        "annotated_measures": f"All **{F['annotated_measures']}** measures",
        "t49_share": f"for **{F['t49_share_rows_matching']}** of {F['stub_rows_main']} rows. **{F['t49_share_rows_other_base']}** special-population rows",
        "t01_cma": f"Table 1 is **{facts['t01_cma_pct_A_to_D']}%** A–D", "t03_cma": f"Table 3 is **{facts['t03_cma_pct_A_to_D']}%** A–D",
        "t04_cma": f"**{facts['t04_cma_pct_F']}%** of estimates are F",
        "t05_cma": f"only **{facts['t05_cma_pct_A_to_D']}%** of estimates are A–D and **{facts['t05_cma_pct_F']}%** are F",
        "t05_industry": f"By industry, **{facts['t05_industry_pct_F']}%** are F",
        "nonprofit_bound": f"total **{F['nonprofit_se_weighted']}** (**{facts['nonprofit_se_pct_of_all']}%** of the All-SMEs count)",
        "portal_bytes": f"**{PORTAL_FILE_BYTES:,} bytes**",
    }
    doc_text = DOC_FILE.read_text(encoding="utf-8")
    if DOC_SECTION_HEADING in doc_text:
        section = doc_text.split(DOC_SECTION_HEADING, 1)[1].split("\n## ", 1)[0]
        missing = {k: v for k, v in doc_claims.items() if v not in section}
        gate("documentation_numbers_reproduced_by_outputs", not missing,
             f"{len(doc_claims) - len(missing)} of {len(doc_claims)} claims found" + (f"; missing {missing}" if missing else ""))
    else:
        gate("documentation_numbers_reproduced_by_outputs", False, "NOT YET DOCUMENTED: Phase C section absent")

    # ---------------------------------------------------
    # 12. Source integrity (after) + summary
    # ---------------------------------------------------
    banner("12. SOURCE INTEGRITY (AFTER)")
    sha_after, stat_after = sha256_file(RAW_FILE), RAW_FILE.stat()
    gate("workbook_sha256_matches_manifest_after", sha_after == m0.sha256)
    gate("workbook_size_and_mtime_unchanged",
         stat_after.st_size == stat_before.st_size and stat_after.st_mtime == stat_before.st_mtime)

    summary = {
        "source": "Survey on Financing and Growth of Small and Medium Enterprises, 2023 - data tables",
        "raw_file": rel, "bytes": stat_before.st_size, "sha256": sha_before,
        "file_modified_local": datetime.fromtimestamp(stat_before.st_mtime).astimezone().isoformat(timespec="seconds"),
        "provenance": {
            "producer": "Statistics Canada, in collaboration with ISED Small Business Branch (cover page)",
            "release_date_cover": release.group(1) if release else None,
            "portal_resource_url": PORTAL_RESOURCE_URL, "portal_file_bytes": PORTAL_FILE_BYTES,
            "manifest_title": m0.table_title, "manifest_reference_period": m0.metadata_reference_period,
            "retrieval_date": m0.retrieval_date,
        },
        "package": {k: v for k, v in pkg.items() if k not in ("sheets", "defined_names")},
        "defined_names_by_kind": dict(Counter(d["kind"] for d in pkg["defined_names"])),
        "sheet_states": dict(states), "formulas": formulas,
        "facts": facts,
        "estimate_types": est.estimate_type.value_counts().to_dict(),
        "value_states": cells.value_state.value_counts().to_dict(),
        "dotdot_by_context": dot_by_ctx,
        "quality_error_ranges": err_ranges.to_dict(orient="records"),
        "number_formats": {f"{u} | {f}": int(n) for (u, f), n in fmts.items()},
        "dimension_counts_main_stub": dim_counts,
        "core_reliability_pct": core_reliability,
        "cma_labels_unmatched_1097": unmatched,
        "relevance_counts": rel_counts,
        "doc_claims": doc_claims,
        "gates": gates,
    }
    summary["all_gates_pass"] = all(gates.values())
    summary["elapsed_seconds"] = round(time.time() - start, 1)
    with open(OUTPUT_DIR / "phase_c_summary.json", "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, default=str, ensure_ascii=False)

    banner("PHASE C COMPLETE")
    failed = [k for k, v in gates.items() if not v]
    print(f"Gates: {len(gates) - len(failed)} of {len(gates)} pass")
    for k in failed:
        print(f"  FAILED: {k}")
    print(f"Outputs: {OUTPUT_DIR}")
    print(f"Elapsed: {summary['elapsed_seconds']}s")
    if failed:
        sys.exit(1)


if __name__ == "__main__":
    main()
